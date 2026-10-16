#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""La moulinette B2B : requalifier et revalider une base de contacts, en local.

Quatre commandes :

  personas    trie une colonne d'intitulés de poste selon une matrice
              (catégorie, séniorité, rôle d'achat déduit).
  b2b         revalide un export (comptes et contacts) selon la moulinette d'une
              mission : règles, liste d'exclusion, doublons, contrôle avant
              lancement (prévol), classeur de contrôle et bilan.
  prevol      dit seulement si la mission peut tourner, contrôle par contrôle.
  empreintes  calcule les empreintes que la mission mémorise (stratégie,
              recherche, filtres, liste d'exclusion).

Exemples :

  python moulinette.py personas contacts.xlsx
  python moulinette.py b2b export.csv --moulinette mission.json --blacklist exclusions.csv
  python moulinette.py b2b export.csv --moulinette mission.json --diagnostic --json

Le fichier de la mission est le même que celui de la page web de la moulinette :
un JSON de configuration (ou un setup complet qui la contient). Rien ne sort de la
machine : aucun appel réseau, aucune donnée envoyée.

Codes de sortie : 0 terminé · 2 prévol bloqué (rien n'est livré) · 1 erreur d'entrée.

Ce moteur est le jumeau exact de celui de la page web (src/lib/segmentation de la
Marketplace ZEROGRAVITY) : mêmes règles, mêmes verdicts, mêmes empreintes. Le banc
`tests/test_moulinette.py` le prouve sur des scénarios rendus par le moteur web.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import json
import math
import re
import sys
import unicodedata
import uuid
from collections import Counter, OrderedDict
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

VERSION = "1.0.0"
ICI = Path(__file__).resolve().parent
MATRICE_DEFAUT = ICI.parent / "matrices" / "matrice-generique.json"

# --------------------------------------------------------------------------- utilitaires « comme le navigateur »


class _Indefini:
    """Valeur absente, distincte de null : une clé indéfinie ne compte pas dans une empreinte."""

    _instance = None

    def __new__(cls):
        if cls._instance is None:
            cls._instance = super().__new__(cls)
        return cls._instance

    def __repr__(self) -> str:
        return "INDEFINI"


INDEFINI = _Indefini()


def nombre_js(n: float | int) -> str:
    """Écrit un nombre comme le navigateur (`String(10.0)` donne « 10 », pas « 10.0 »)."""
    if isinstance(n, bool):
        return "true" if n else "false"
    if isinstance(n, int):
        return str(n)
    if math.isnan(n):
        return "NaN"
    if math.isinf(n):
        return "Infinity" if n > 0 else "-Infinity"
    if n.is_integer() and abs(n) < 1e21:
        return str(int(n))
    texte = repr(n)
    if "e" in texte:
        mantisse, exposant = texte.split("e")
        valeur = int(exposant)
        texte = f"{mantisse}e{'+' if valeur > 0 else '-'}{abs(valeur)}"
    return texte


def texte_js(valeur: Any) -> str:
    """`String(valeur ?? "")` : None et l'absence donnent une chaîne vide."""
    if valeur is None or valeur is INDEFINI:
        return ""
    if isinstance(valeur, bool):
        return "true" if valeur else "false"
    if isinstance(valeur, (int, float)):
        return nombre_js(valeur)
    if isinstance(valeur, (dt.datetime, dt.date)):
        return valeur.isoformat()
    return str(valeur)


def chaine_js(valeur: Any) -> str:
    """`String(valeur)` strict : null donne « null »."""
    if valeur is None:
        return "null"
    return texte_js(valeur)


_ESPACES = "\t\n\v\f\r                  　﻿"


def rogner(texte: str) -> str:
    """`trim()` du navigateur (il retire aussi l'indicateur d'ordre des octets)."""
    return texte.strip(_ESPACES)


def est_nombre(valeur: Any) -> bool:
    return isinstance(valeur, (int, float)) and not isinstance(valeur, bool)


def est_fini(valeur: Any) -> bool:
    """`Number.isFinite` : un vrai nombre, jamais une chaîne qui en a l'air."""
    return est_nombre(valeur) and math.isfinite(valeur)


def vrai(valeur: Any) -> bool:
    """Vérité au sens du navigateur (0, "", None et l'absence sont faux)."""
    if valeur is None or valeur is INDEFINI:
        return False
    if isinstance(valeur, float) and math.isnan(valeur):
        return False
    if isinstance(valeur, (list, dict)):
        return True
    return bool(valeur)


def json_js(valeur: Any) -> str:
    """`JSON.stringify` d'une valeur simple."""
    if valeur is None or valeur is INDEFINI:
        return "null"
    if isinstance(valeur, bool):
        return "true" if valeur else "false"
    if est_nombre(valeur):
        return nombre_js(valeur) if math.isfinite(valeur) else "null"
    if isinstance(valeur, str):
        return json.dumps(valeur, ensure_ascii=False)
    return json.dumps(valeur, ensure_ascii=False, separators=(",", ":"))


def _cle_locale(cle: str) -> tuple:
    """Ordre des clés de `localeCompare` : sans casse d'abord, minuscule avant majuscule ensuite."""
    return (cle.casefold(), cle.swapcase())


def serialiser_stable(valeur: Any) -> str:
    """Sérialisation stable au sens de JSON (clé indéfinie omise). Jumelle de `serialiserStable`."""
    if isinstance(valeur, list):
        return "[" + ",".join("null" if item is INDEFINI else serialiser_stable(item) for item in valeur) + "]"
    if isinstance(valeur, dict):
        entrees = sorted(((cle, contenu) for cle, contenu in valeur.items() if contenu is not INDEFINI), key=lambda kv: _cle_locale(kv[0]))
        return "{" + ",".join(f"{json.dumps(cle, ensure_ascii=False)}:{serialiser_stable(contenu)}" for cle, contenu in entrees) + "}"
    return json_js(valeur)


def fnv1a(texte: str) -> str:
    """FNV-1a 32 bits sur les unités UTF-16, comme `charCodeAt` du navigateur."""
    empreinte = 0x811C9DC5
    octets = texte.encode("utf-16-le", "surrogatepass")
    for index in range(0, len(octets), 2):
        empreinte ^= octets[index] | (octets[index + 1] << 8)
        empreinte = (empreinte * 0x01000193) & 0xFFFFFFFF
    return f"{empreinte:08x}"


def empreinte_texte(valeur: Any) -> str:
    return f"fnv1a-{fnv1a(serialiser_stable(valeur))}"


def _id_local() -> str:
    return str(uuid.uuid4())


# --------------------------------------------------------------------------- normalisations

DOMAINES_EMAIL_PERSONNELS = {
    "gmail.com", "googlemail.com", "hotmail.com", "hotmail.fr", "outlook.com", "outlook.fr", "live.com",
    "live.fr", "yahoo.com", "yahoo.fr", "icloud.com", "orange.fr", "free.fr", "laposte.net", "proton.me",
    "protonmail.com",
}
FORMES_JURIDIQUES = {"sas", "sasu", "sarl", "eurl", "sa", "sca", "scop", "gie", "ei", "ltd", "limited", "inc", "llc"}


# `\p{Diacritic}` du navigateur : les marques combinantes, et aussi les accents isolés (^ ` ¨ ¯ ´ · ¸)
# et les lettres modificatives. Quelques marques combinantes latines n'en sont pas.
_ACCENTS_ISOLES = {0x5E, 0x60, 0xA8, 0xAF, 0xB4, 0xB7, 0xB8}
_MARQUES_NON_DIACRITIQUES = {0x34F, *range(0x358, 0x35D), *range(0x363, 0x370)}


def _diacritique(caractere: str) -> bool:
    code = ord(caractere)
    if code in _ACCENTS_ISOLES or 0x2B0 <= code <= 0x2FF:
        return True
    return unicodedata.category(caractere) == "Mn" and code not in _MARQUES_NON_DIACRITIQUES


def sans_diacritiques(texte: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", texte) if not _diacritique(c))


def normaliser_texte(valeur: Any) -> str:
    """Minuscules, sans accents, ponctuation ramenée à une espace."""
    if valeur is None or valeur is INDEFINI:
        return ""
    texte = sans_diacritiques(texte_js(valeur)).lower().replace("’", " ").replace("'", " ")
    return re.sub(r"[^a-z0-9]+", " ", texte).strip(" ")


def normaliser_nom_entreprise(valeur: Any) -> str:
    texte = normaliser_texte(valeur)
    texte = re.sub(r"\b(s a s u|s a s|s a r l|e u r l|s a|l t d|l l c)\b", lambda m: m.group(0).replace(" ", ""), texte)
    return " ".join(mot for mot in texte.split(" ") if mot not in FORMES_JURIDIQUES).strip(" ")


_HOTE_INTERDIT = re.compile(r"[\x00-\x20#/:<>?@\[\\\]^|%\x7f\"]")


def _encoder_chemin(chemin: str) -> str:
    """Encodage du chemin d'une URL, comme le navigateur (`url.pathname`)."""
    sortie = []
    for caractere in chemin.replace("\\", "/"):
        code = ord(caractere)
        if code <= 0x20 or code >= 0x7F or caractere in "\"<>`{}":
            sortie.extend(f"%{octet:02X}" for octet in caractere.encode("utf-8", "surrogatepass"))
        else:
            sortie.append(caractere)
    segments: list[str] = []
    morceaux = "".join(sortie).split("/")
    for position, morceau in enumerate(morceaux[1:] if morceaux and morceaux[0] == "" else morceaux):
        dernier = position == len(morceaux) - 2
        bas = morceau.lower()
        if bas in {".", "%2e"}:
            if dernier:
                segments.append("")
            continue
        if bas in {"..", ".%2e", "%2e.", "%2e%2e"}:
            if segments:
                segments.pop()
            if dernier:
                segments.append("")
            continue
        segments.append(morceau)
    return "/" + "/".join(segments)


def _analyser_url(brut: str) -> tuple[str, str] | None:
    """(hôte, chemin) d'une adresse, ou None si le navigateur la refuserait."""
    texte = brut if re.match(r"^https?://", brut, re.I) else f"https://{brut}"
    texte = re.sub(r"[\t\n\r]", "", texte)
    try:
        morceaux = urlsplit(texte)
        hote = morceaux.hostname
    except ValueError:
        return None
    if not hote or _HOTE_INTERDIT.search(hote):
        return None
    try:
        hote = ".".join(etiquette.encode("idna").decode("ascii") if etiquette else "" for etiquette in hote.split("."))
    except UnicodeError:
        return None
    hote = _hote_ipv4(hote.lower())
    if hote is None:
        return None
    return hote, _encoder_chemin(morceaux.path or "/")


def _nombre_ipv4(partie: str) -> int | None:
    base = 10
    if partie[:2] in ("0x", "0X"):
        base, partie = 16, partie[2:]
    elif len(partie) > 1 and partie.startswith("0"):
        base, partie = 8, partie[1:]
    if partie == "":
        return 0
    try:
        return int(partie, base) if re.fullmatch(r"[0-9a-fA-F]+", partie) else None
    except ValueError:
        return None


def _hote_ipv4(hote: str) -> str | None:
    """Un hôte qui finit par un nombre est une adresse IPv4 pour le navigateur (« 1.2.3 » devient 1.2.0.3)."""
    parties = hote.split(".")
    if parties[-1] == "" and len(parties) > 1:
        parties = parties[:-1]
    dernier = parties[-1]
    if not (re.fullmatch(r"[0-9]+", dernier) or (dernier[:2] in ("0x", "0X") and re.fullmatch(r"[0-9a-fA-F]*", dernier[2:]))):
        return hote
    if len(parties) > 4 or any(p == "" for p in parties):
        return None
    nombres = [_nombre_ipv4(p) for p in parties]
    if any(n is None for n in nombres) or any(n > 255 for n in nombres[:-1]) or nombres[-1] >= 256 ** (5 - len(nombres)):
        return None
    valeur = nombres[-1] + sum(n * 256 ** (3 - i) for i, n in enumerate(nombres[:-1]))
    return ".".join(str((valeur >> decalage) & 255) for decalage in (24, 16, 8, 0))


def normaliser_url_linkedin(valeur: Any) -> str:
    brut = rogner(texte_js(valeur))
    if not brut:
        return ""
    analyse = _analyser_url(brut)
    if not analyse:
        return ""
    hote = re.sub(r"^www\.", "", analyse[0])
    if not hote.endswith("linkedin.com"):
        return ""
    return "https://www.linkedin.com" + re.sub(r"/+$", "", analyse[1]).lower()


def normaliser_domaine(valeur: Any) -> str:
    brut = rogner(texte_js(valeur)).lower()
    if not brut:
        return ""
    analyse = _analyser_url(brut)
    if analyse:
        return re.sub(r"\.$", "", re.sub(r"^www\.", "", analyse[0]))
    sans_protocole = re.sub(r"^www\.", "", re.sub(r"^https?://", "", brut))
    return re.sub(r"\.$", "", re.split(r"[/?#]", sans_protocole)[0])


def normaliser_siren(valeur: Any) -> str:
    chiffres = re.sub(r"[^0-9]", "", texte_js(valeur))
    return chiffres if len(chiffres) == 9 else ""


def normaliser_categorie_linkedin(valeur: Any) -> str:
    bornes = re.search(r"([0-9]{1,5})[^0-9]+([0-9]{1,5})", normaliser_texte(valeur))
    return f"{int(bornes.group(1))}-{int(bornes.group(2))}" if bornes else ""


def lire_nombre(valeur: Any) -> float | int | None:
    if est_nombre(valeur):
        return valeur if math.isfinite(valeur) else None
    nettoyee = re.sub(r"\s", "", texte_js(valeur)).replace(",", ".", 1)
    morceaux = re.findall(r"-?[0-9]+(?:\.[0-9]+)?", nettoyee)
    if not morceaux:
        return None
    try:
        resultat = float("".join(morceaux))
    except ValueError:
        return None
    if not math.isfinite(resultat):
        return None
    return int(resultat) if resultat.is_integer() else resultat


def normaliser_email(valeur: Any) -> str:
    return rogner(texte_js(valeur)).lower()


def email_est_professionnel(valeur: Any) -> bool:
    email = normaliser_email(valeur)
    return bool(re.match(r"^[^\s@]+@[^\s@]+\.[^\s@]+$", email)) and email.split("@")[1] not in DOMAINES_EMAIL_PERSONNELS


def normaliser_mobile_francais(valeur: Any) -> str:
    chiffres = re.sub(r"[^0-9]", "", texte_js(valeur))
    if chiffres.startswith("0033"):
        chiffres = chiffres[4:]
    elif chiffres.startswith("33"):
        chiffres = chiffres[2:]
    if chiffres.startswith("0"):
        chiffres = chiffres[1:]
    return f"0{chiffres}" if re.fullmatch(r"[67][0-9]{8}", chiffres) else ""


# --------------------------------------------------------------------------- critères (jumeau de regles.ts)

ENTETES_CANDIDATS: "OrderedDict[str, list[str]]" = OrderedDict([
    ("nomComplet", ["nom complet", "full name", "nom", "contact"]),
    ("entreprise", ["entreprise", "company", "company name", "societe", "account"]),
    ("poste", ["poste", "job title", "title", "fonction"]),
    ("linkedinProfil", ["linkedin profil", "linkedin profile", "linkedin url", "person linkedin url"]),
    ("linkedinEntreprise", ["linkedin entreprise", "company linkedin", "company linkedin url"]),
    ("categorieLinkedIn", ["categorie linkedin", "categorie linkedin 11 50", "linkedin company size", "company size", "taille linkedin"]),
    ("employesVisibles", ["nombre d employes visibles", "employes visibles", "employees on linkedin", "linkedin employees"]),
    ("pays", ["pays", "country", "company country"]),
    ("activite", ["activite", "activity", "industry", "industrie", "description"]),
    ("tagsFournisseur", ["tags fournisseur", "tags pharow", "tags", "source tags", "provider tags"]),
    ("trancheEffectifFournisseur", ["tranche effectif fournisseur", "tranche effectif pharow", "effectif france pharow", "company headcount range"]),
    ("qualificationActivite", ["qualification activite", "qualification metier qa", "qa activite", "business model qa"]),
    ("siteEntreprise", ["site entreprise", "company website", "website", "site", "domain", "domaine"]),
    ("siren", ["siren", "numero siren", "company siren"]),
    ("emailProfessionnel", ["email professionnel", "work email", "professional email", "email"]),
    ("sourceEmail", ["source email", "email source", "provenance email", "provenance de l email"]),
    ("mobile", ["mobile", "mobile obligatoire 06 07", "telephone mobile", "phone", "telephone", "numero de telephone"]),
    ("sourceMobile", ["source mobile", "phone source", "mobile source", "provenance mobile", "provenance du mobile", "provenance professionnelle du mobile"]),
    ("effectifLegal", ["effectif legal", "effectif france", "legal headcount"]),
    ("chiffreAffaires", ["chiffre d affaires", "chiffre affaires", "ca", "revenue"]),
    ("codeNaf", ["code naf", "naf", "ape", "code ape"]),
    ("statutJuridique", ["statut juridique", "forme juridique", "legal form"]),
    ("technologies", ["technologies", "tech stack", "stack"]),
    ("departement", ["departement", "department", "fonction departement"]),
    ("niveauHierarchique", ["niveau hierarchique", "seniority", "seniorite"]),
    ("anciennetePoste", ["anciennete poste", "anciennete dans le poste", "tenure"]),
])
ENTETES_BLACKLIST: "OrderedDict[str, list[str]]" = OrderedDict(
    (champ, ENTETES_CANDIDATS[champ]) for champ in ("entreprise", "linkedinEntreprise", "siteEntreprise", "siren")
)
CHAMPS_NUMERIQUES = {"employesVisibles", "effectifLegal", "chiffreAffaires", "anciennetePoste"}
CHAMPS_CONTACT = {"nomComplet", "poste", "linkedinProfil", "emailProfessionnel", "mobile", "sourceEmail", "sourceMobile", "departement", "niveauHierarchique", "anciennetePoste"}

# id -> (champ, libellé, portée)
DEFINITIONS: dict[str, tuple[str | None, str, str]] = {
    "pays": ("pays", "Pays", "entreprise"),
    "activite": ("activite", "Activité / expressions", "entreprise"),
    "tags_source": ("tagsFournisseur", "Tag du fournisseur", "entreprise"),
    "tranche_effectif_source": ("trancheEffectifFournisseur", "Tranche d'effectif du fournisseur", "entreprise"),
    "qualification_activite": ("qualificationActivite", "Qualification métier QA", "entreprise"),
    "categorie_linkedin": ("categorieLinkedIn", "Catégorie LinkedIn", "entreprise"),
    "employes_visibles": ("employesVisibles", "Employés visibles LinkedIn", "entreprise"),
    "effectif_legal": ("effectifLegal", "Effectif légal", "entreprise"),
    "chiffre_affaires": ("chiffreAffaires", "Chiffre d'affaires", "entreprise"),
    "code_naf": ("codeNaf", "Code NAF", "entreprise"),
    "statut_juridique": ("statutJuridique", "Statut juridique", "entreprise"),
    "technologies": ("technologies", "Technologies", "entreprise"),
    "poste": ("poste", "Poste du décideur", "contact"),
    "departement": ("departement", "Département", "contact"),
    "niveau_hierarchique": ("niveauHierarchique", "Niveau hiérarchique", "contact"),
    "email_pro": ("emailProfessionnel", "Email professionnel", "contact"),
    "source_email": ("sourceEmail", "Provenance de l'email", "contact"),
    "mobile_fr": ("mobile", "Mobile français 06/07", "contact"),
    "source_mobile": ("sourceMobile", "Provenance professionnelle du mobile", "contact"),
    "blacklist": (None, "Entreprise absente de la blacklist", "exclusion"),
    "doublon": (None, "Contact unique", "qualite"),
}

COLONNES_B2B_STANDARD = [
    {"champ": "nomComplet", "libelle": "Nom complet"},
    {"champ": "entreprise", "libelle": "Entreprise"},
    {"champ": "poste", "libelle": "Poste"},
    {"champ": "linkedinProfil", "libelle": "LinkedIn profil"},
    {"champ": "categorieLinkedIn", "libelle": "Catégorie LinkedIn"},
    {"champ": "employesVisibles", "libelle": "Nombre d'employés visibles"},
    {"champ": "emailProfessionnel", "libelle": "Email professionnel"},
    {"champ": "mobile", "libelle": "Mobile professionnel"},
]


# --------------------------------------------------------------------------- lecture d'une mission (jumeau de migrerConfigurationB2B)


def _regle(critere_id: str, operateur: str, valeur: Any, importance: str = "bloquant", si_absent: str = "rejeter") -> dict:
    return {"id": _id_local(), "critereId": critere_id, "operateur": operateur, "valeur": valeur, "importance": importance, "siAbsent": si_absent, "active": True}


def _defaut(valeur: Any, defaut: Any) -> Any:
    """L'opérateur `??` : None et l'absence prennent la valeur par défaut."""
    return defaut if valeur is None or valeur is INDEFINI else valeur


def migrer_filtres_source(filtres: Any) -> list[dict]:
    if not isinstance(filtres, list):
        return []
    sortie: list[dict] = []
    for filtre in filtres:
        if isinstance(filtre, str):
            libelle, _, reste = filtre.partition(":")
            sortie.append({
                "id": _id_local(), "recherche": "Recherche historique", "etape": "comptes",
                "libelle": rogner(libelle) or "Filtre historique", "valeurs": [rogner(reste) or "À confirmer"],
                "logique": "ET", "preuve": "configuration", "obligatoire": True,
                "note": "Filtre migré : relier à une règle avant le prochain run si ce critère porte sur chaque ligne.",
            })
            continue
        if not isinstance(filtre, dict):
            continue
        preuve = filtre.get("preuve")
        valeurs = filtre.get("valeurs")
        recu = {
            "id": filtre.get("id") or _id_local(),
            "recherche": rogner(texte_js(filtre.get("recherche"))) or "Recherche sans nom",
            "etape": "contacts" if filtre.get("etape") == "contacts" else "comptes",
            "libelle": rogner(texte_js(filtre.get("libelle"))) or "Filtre sans nom",
            "valeurs": [v for v in (chaine_js(x) for x in valeurs) if v] if isinstance(valeurs, list) else ["À confirmer"],
            "logique": filtre.get("logique") if filtre.get("logique") in ("ET", "OU", "EXCLURE") else "ET",
            "preuve": preuve if preuve in ("export", "configuration", "calcul", "qa_manuelle") else "configuration",
            "regleId": filtre.get("regleId", INDEFINI),
            "obligatoire": filtre.get("obligatoire") is not False,
            "note": filtre.get("note", INDEFINI),
        }
        sortie.append(recu)
    return sortie


def migrer_strategie(strategie: Any) -> dict:
    if not isinstance(strategie, dict):
        return {"lignes": [], "limites": ""}
    lignes = strategie.get("lignes")
    return {
        "lignes": [{**ligne, "id": ligne.get("id") or _id_local(), "obligatoire": ligne.get("obligatoire") is not False} for ligne in lignes if isinstance(ligne, dict)] if isinstance(lignes, list) else [],
        "limites": _defaut(strategie.get("limites"), ""),
        "valideeLe": strategie.get("valideeLe", INDEFINI),
        "empreinteValidee": strategie.get("empreinteValidee", INDEFINI),
    }


def _pharow(source: dict) -> dict:
    pharow = source.get("pharow") if isinstance(source.get("pharow"), dict) else {}
    return {
        "nomRecherche": "", "nomListeExclusion": "", **pharow,
        "filtres": migrer_filtres_source(pharow.get("filtres")),
        "echantillonTemoins": _defaut(pharow.get("echantillonTemoins"), []),
    }


def migrer_configuration_b2b(source: dict) -> dict:
    blacklist_defaut = {"obligatoire": True, "nom": "Liste d'exclusion", "typeSource": "fichier"}
    if source.get("schemaVersion") == 2 and isinstance(source.get("regles"), list):
        return {
            **source, "mode": "revalidation_b2b", "schemaVersion": 2, "regles": source["regles"],
            "champsObligatoires": _defaut(source.get("champsObligatoires"), []),
            "colonnesFinales": _defaut(source.get("colonnesFinales"), [dict(c) for c in COLONNES_B2B_STANDARD]),
            "mappingCandidats": _defaut(source.get("mappingCandidats"), {}),
            "mappingBlacklist": _defaut(source.get("mappingBlacklist"), {}),
            "blacklist": _defaut(source.get("blacklist"), blacklist_defaut),
            "strategie": migrer_strategie(source.get("strategie")),
            "pharow": _pharow(source),
        }
    ancien = source.get("criteres") if isinstance(source.get("criteres"), dict) else {}
    regles: list[dict] = []
    if ancien.get("paysAcceptes"):
        regles.append(_regle("pays", "est_dans", ancien["paysAcceptes"]))
    if ancien.get("motsClesActivite"):
        regles.append(_regle("activite", "contient_un", ancien["motsClesActivite"]))
    if ancien.get("categoriesLinkedIn"):
        regles.append(_regle("categorie_linkedin", "est_dans", ancien["categoriesLinkedIn"]))
    if est_fini(ancien.get("employesVisiblesMin")) and est_fini(ancien.get("employesVisiblesMax")):
        regles.append(_regle("employes_visibles", "entre", {"min": ancien["employesVisiblesMin"], "max": ancien["employesVisiblesMax"]}))
    if ancien.get("postesAcceptes"):
        regles.append(_regle("poste", "est_dans", ancien["postesAcceptes"]))
    if ancien.get("emailProfessionnelObligatoire"):
        regles.append(_regle("email_pro", "email_professionnel", True))
    if ancien.get("sourceEmailObligatoire"):
        regles.append(_regle("source_email", "present", True, "verification", "verifier"))
    if ancien.get("mobileProfessionnelObligatoire"):
        regles.append(_regle("mobile_fr", "mobile_francais", True))
    if ancien.get("sourceMobileObligatoire"):
        regles.append(_regle("source_mobile", "provenance_professionnelle", True, "verification", "verifier"))
    regles += [_regle("blacklist", "pas_blackliste", True), _regle("doublon", "unique", True)]
    return {
        "mode": "revalidation_b2b", "schemaVersion": 2, "regles": regles,
        "champsObligatoires": _defaut(ancien.get("champsObligatoires"), []),
        "colonnesFinales": _defaut(ancien.get("colonnesFinales"), [dict(c) for c in COLONNES_B2B_STANDARD]),
        "mappingCandidats": _defaut(source.get("mappingCandidats"), {}),
        "mappingBlacklist": _defaut(source.get("mappingBlacklist"), {}),
        "blacklist": _defaut(source.get("blacklist"), blacklist_defaut),
        "strategie": {"lignes": [], "limites": ""},
        "pharow": _pharow(source),
    }


def charger_mission(chemin: Path) -> tuple[dict, dict]:
    """Rend (identité de la mission, configuration). Accepte un setup, une ligne de base ou une configuration nue."""
    try:
        brut = json.loads(chemin.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as erreur:
        raise ErreurEntree(f"Mission illisible ({chemin}) : {erreur}") from erreur
    if not isinstance(brut, dict):
        raise ErreurEntree(f"Mission illisible ({chemin}) : un objet JSON est attendu.")
    configuration = brut.get("configuration") if isinstance(brut.get("configuration"), dict) else brut
    identite = {cle: brut.get(cle) for cle in ("id", "client", "mission", "description") if brut.get(cle)}
    if configuration.get("mode") == "persona" or "axes" in configuration or "matrice" in configuration:
        return identite, {"mode": "persona", "matrice": configuration.get("matrice", configuration), "multiCategories": bool(configuration.get("multiCategories"))}
    return identite, migrer_configuration_b2b(configuration)


# --------------------------------------------------------------------------- empreintes et prévol (jumeau de setups.ts et b2b.ts)


def empreinte_strategie(strategie: dict) -> str:
    return empreinte_texte({"lignes": strategie.get("lignes", []), "limites": rogner(texte_js(strategie.get("limites")))})


def empreinte_filtres(filtres: list) -> str:
    return empreinte_texte(filtres)


def empreinte_recherche(configuration: dict) -> str:
    pharow = configuration["pharow"]
    return empreinte_texte({
        "strategie": empreinte_strategie(configuration["strategie"]),
        "filtres": empreinte_filtres(pharow["filtres"]),
        "nomRecherche": rogner(texte_js(pharow.get("nomRecherche"))),
        "nomListeExclusion": rogner(texte_js(pharow.get("nomListeExclusion"))),
        "volumeResultats": pharow.get("volumeResultats", INDEFINI),
        "echantillonTemoins": [rogner(texte_js(item)) for item in pharow.get("echantillonTemoins", []) if rogner(texte_js(item))],
        "captureNom": pharow.get("captureNom", INDEFINI),
        "captureEmpreinte": pharow.get("captureEmpreinte", INDEFINI),
    })


PARENTS_GENERIQUES = {"informatique", "it", "technologie", "technologies", "services", "logiciel", "logiciels"}
# `\W` du navigateur = tout sauf [A-Za-z0-9_] : « é » y compte comme séparateur.
TERMES_ESN_TROP_LARGES = re.compile(r"(^|[^A-Za-z0-9_])(informatique|it|software|data|cloud|développement|conseil tech)([^A-Za-z0-9_]|$)", re.I)


def erreurs_strategie(strategie: dict) -> list[str]:
    lignes = strategie.get("lignes", [])
    erreurs: list[str] = []
    if not lignes:
        erreurs.append("La stratégie ne contient aucun critère.")
    if not rogner(texte_js(strategie.get("limites"))):
        erreurs.append("Les limites connues des filtres ne sont pas décrites.")
    champ = lambda ligne, cle: rogner(texte_js(ligne.get(cle)))  # noqa: E731
    if any(vrai(l.get("obligatoire")) and (not champ(l, "critereClient") or not champ(l, "parametrage") or not champ(l, "controleAval")) for l in lignes):
        erreurs.append("Chaque critère obligatoire doit préciser le besoin, le paramétrage exact et le contrôle aval.")
    principales = [l for l in lignes if vrai(l.get("obligatoire")) and l.get("categorie") == "type_entreprise" and l.get("logique") != "EXCLURE" and l.get("precision") == "granulaire"]
    if not principales:
        erreurs.append("Le type d'entreprise exige au moins une voie principale granulaire.")

    def uniquement_generique(parametrage: str) -> bool:
        valeurs = [rogner(v).lower() for v in re.split(r"[;,/|\n]+", parametrage.split(":")[-1])]
        return all(v in PARENTS_GENERIQUES for v in valeurs if v)

    if any(uniquement_generique(texte_js(l.get("parametrage"))) for l in principales):
        erreurs.append("Un secteur parent générique ne constitue pas une stratégie sectorielle granulaire.")
    if any(";" in texte_js(l.get("parametrage")) or "," in texte_js(l.get("parametrage")) for l in principales):
        erreurs.append("Les critères d'activité doivent conserver leur logique booléenne : une condition par ligne avec ET, OU ou EXCLURE.")
    if any(re.search(r"esn", texte_js(l.get("critereClient")), re.I) and TERMES_ESN_TROP_LARGES.search(texte_js(l.get("parametrage"))) for l in principales):
        erreurs.append("Une recherche ESN ne peut pas accepter un terme IT générique comme condition sectorielle.")
    return erreurs


def strategie_validee(strategie: dict) -> bool:
    return not erreurs_strategie(strategie) and vrai(strategie.get("valideeLe")) and strategie.get("empreinteValidee") == empreinte_strategie(strategie)


def recherche_validee(configuration: dict) -> bool:
    pharow = configuration["pharow"]
    temoins = [item for item in pharow.get("echantillonTemoins", []) if rogner(texte_js(item))]
    return (
        strategie_validee(configuration["strategie"])
        and bool(rogner(texte_js(pharow.get("nomRecherche"))) and rogner(texte_js(pharow.get("nomListeExclusion"))))
        and len(pharow["filtres"]) > 0
        and all(rogner(f["recherche"]) and rogner(f["libelle"]) and len(f["valeurs"]) > 0 for f in pharow["filtres"])
        and est_fini(pharow.get("volumeResultats")) and pharow["volumeResultats"] > 0
        and 5 <= len(temoins) <= 10
        and vrai(pharow.get("captureNom")) and vrai(pharow.get("captureEmpreinte")) and vrai(pharow.get("rechercheValideeLe"))
        and pharow.get("rechercheValideeEmpreinte") == empreinte_recherche(configuration)
    )


def _index_blacklist(lignes: list[dict], mapping: dict) -> dict[str, set]:
    index = {"linkedin": set(), "domaines": set(), "sirens": set(), "noms": set()}
    for ligne in lignes:
        lire = lambda champ: ligne.get(mapping[champ]) if vrai(mapping.get(champ)) else None  # noqa: E731
        for cle, valeur in (
            ("linkedin", normaliser_url_linkedin(lire("linkedinEntreprise"))),
            ("domaines", normaliser_domaine(lire("siteEntreprise"))),
            ("sirens", normaliser_siren(lire("siren"))),
            ("noms", normaliser_nom_entreprise(lire("entreprise"))),
        ):
            if valeur:
                index[cle].add(valeur)
    return index


def empreinte_blacklist(lignes: list[dict], mapping: dict) -> str:
    index = _index_blacklist(lignes, mapping)
    identifiants = sorted(
        [f"linkedin:{v}" for v in index["linkedin"]] + [f"domaine:{v}" for v in index["domaines"]]
        + [f"siren:{v}" for v in index["sirens"]] + [f"nom:{v}" for v in index["noms"]],
        key=lambda texte: texte.encode("utf-16-be", "surrogatepass"),
    )
    return f"fnv1a-{fnv1a(chr(10).join(identifiants))}"


def controler_prevol(configuration: dict, candidats_charges: int, blacklist_chargee: int, empreinte_bl: str | None) -> list[dict]:
    regles = [r for r in configuration["regles"] if vrai(r.get("active"))]
    pharow = configuration["pharow"]
    mapping = configuration["mappingCandidats"]
    reference = configuration["blacklist"]
    ids_valides = all(r.get("critereId") in DEFINITIONS for r in regles)
    blacklist_requise = vrai(reference.get("obligatoire")) or any(r.get("operateur") == "pas_blackliste" for r in regles)

    def intervalle_valide(regle: dict) -> bool:
        valeur = regle.get("valeur")
        return isinstance(valeur, dict) and est_fini(valeur.get("min")) and est_fini(valeur.get("max")) and valeur["min"] <= valeur["max"]

    intervalles = all(intervalle_valide(r) for r in regles if r.get("operateur") == "entre")
    listes = all(isinstance(r.get("valeur"), list) and len(r["valeur"]) > 0 for r in regles if r.get("operateur") in ("est_dans", "contient_un", "ne_contient_aucun"))
    champs_regles = [DEFINITIONS[r["critereId"]][0] for r in regles if r.get("critereId") in DEFINITIONS and DEFINITIONS[r["critereId"]][0]]
    mappings_complets = all(vrai(mapping.get(champ)) for champ in set(configuration["champsObligatoires"]) | set(champs_regles))
    mapping_blacklist = any(vrai(v) for v in configuration["mappingBlacklist"].values())
    volume = reference.get("volumeAttendu", INDEFINI)
    volume_conforme = est_nombre(volume) and volume == blacklist_chargee
    empreinte_conforme = vrai(reference.get("empreinte")) and vrai(empreinte_bl) and reference.get("empreinte") == empreinte_bl
    filtres = pharow["filtres"]
    filtres_valides = len(filtres) > 0 and all(vrai(f.get("id")) and rogner(f["recherche"]) and rogner(f["libelle"]) and len(f["valeurs"]) > 0 for f in filtres)
    export_conforme = vrai(pharow.get("exportBrutNom")) and vrai(pharow.get("exportCaptureLe")) and pharow.get("exportBrutVolume", INDEFINI) == candidats_charges and est_nombre(pharow.get("exportBrutVolume")) and candidats_charges > 0
    saisis, capture = pharow.get("filtresSaisisApresExtractionLe"), pharow.get("exportCaptureLe")
    saisis_apres = vrai(saisis) and vrai(capture) and texte_js(saisis) >= texte_js(capture)
    empreinte_filtres_conforme = vrai(pharow.get("filtresConfirmesEmpreinte")) and pharow.get("filtresConfirmesEmpreinte") == empreinte_filtres(filtres)

    def couvert(filtre: dict) -> bool:
        if filtre["preuve"] == "configuration":
            return True
        liee = next((r for r in regles if r.get("id") == filtre.get("regleId") and filtre.get("regleId") is not INDEFINI), None)
        if not liee:
            return False
        if filtre["preuve"] == "calcul":
            return True
        champ = DEFINITIONS.get(liee.get("critereId"), (None,))[0]
        return bool(champ and vrai(mapping.get(champ)))

    couverts = all(couvert(f) for f in filtres if vrai(f.get("obligatoire")))
    controle = lambda ident, libelle, valide, bloquant=True: {"id": ident, "libelle": libelle, "valide": bool(valide), "bloquant": bool(bloquant)}  # noqa: E731
    return [
        controle("strategie", "Stratégie critère par critère validée et inchangée", strategie_validee(configuration["strategie"])),
        controle("recherche_validee", "Recherche réelle validée avec capture, volume, échantillon et limites", recherche_validee(configuration)),
        controle("configuration", "Règles, colonnes finales et mapping candidats complets", len(regles) > 0 and ids_valides and intervalles and listes and len(configuration["colonnesFinales"]) > 0 and mappings_complets),
        controle("blacklist", "Blacklist chargée et mapping d'exclusion défini", blacklist_chargee > 0 and mapping_blacklist, blacklist_requise),
        controle("empreinte", "Volume et empreinte de blacklist cohérents", blacklist_chargee > 0 and volume_conforme and empreinte_conforme, blacklist_requise),
        controle("liste_pharow", "Bonne recherche et bonne liste d'exclusion Pharow renseignées", rogner(texte_js(pharow.get("nomRecherche"))) and rogner(texte_js(pharow.get("nomListeExclusion")))),
        controle("exclusion_pharow", "Exclusion confirmée active dans Pharow", vrai(pharow.get("exclusionConfirmeeLe"))),
        controle("export_source", "Export brut conservé et identifié dans le navigateur", export_conforme),
        controle("filtres", "Filtres définitifs recopiés après l'extraction, confirmés et inchangés", filtres_valides and saisis_apres and vrai(pharow.get("filtresConfirmesLe")) and empreinte_filtres_conforme),
        controle("couverture_filtres", "Chaque filtre obligatoire est relié à une règle et à sa preuve", couverts),
        controle("candidats", "Fichier candidats prêt", candidats_charges > 0),
    ]


def prevol_autorise(controles: list[dict]) -> bool:
    return all(not c["bloquant"] or c["valide"] for c in controles)


# --------------------------------------------------------------------------- validation B2B (jumeau de validerExportB2B)


def deviner_mapping(entetes: list[str], synonymes: "OrderedDict[str, list[str]]") -> dict[str, str]:
    normalisees = {normaliser_texte(entete): entete for entete in entetes}
    resultat: dict[str, str] = {}
    for champ, noms in synonymes.items():
        exact = next((nom for nom in map(normaliser_texte, noms) if nom in normalisees), None)
        if exact is not None:
            resultat[champ] = normalisees[exact]
    return resultat


def completer_mapping(memorise: dict, entetes: list[str], synonymes) -> dict:
    """Comme la page : le mapping deviné, recouvert par le mapping mémorisé quand sa colonne existe."""
    return {**deviner_mapping(entetes, synonymes), **{champ: colonne for champ, colonne in memorise.items() if vrai(colonne) and colonne in entetes}}


def _normaliser_champ(champ: str, brut: Any) -> Any:
    if champ in ("linkedinProfil", "linkedinEntreprise"):
        return normaliser_url_linkedin(brut)
    if champ == "siteEntreprise":
        return normaliser_domaine(brut)
    if champ == "siren":
        return normaliser_siren(brut)
    if champ == "emailProfessionnel":
        return normaliser_email(brut)
    if champ == "mobile":
        return normaliser_mobile_francais(brut)
    if champ in CHAMPS_NUMERIQUES:
        nombre = lire_nombre(brut)
        return "" if nombre is None else nombre
    if champ == "categorieLinkedIn":
        return normaliser_categorie_linkedin(brut)
    return rogner(texte_js(brut))


def valeurs_normalisees(ligne: dict, mapping: dict) -> dict:
    sortie = {}
    for champ in ENTETES_CANDIDATS:
        colonne = mapping.get(champ)
        brut = ligne.get(colonne, INDEFINI) if vrai(colonne) else INDEFINI
        if brut is INDEFINI or brut is None or rogner(texte_js(brut)) == "":
            continue
        normalisee = _normaliser_champ(champ, brut)
        if normalisee != "" or est_nombre(normalisee):
            sortie[champ] = normalisee
    return sortie


def _correspondance_blacklist(valeurs: dict, index: dict) -> str | None:
    linkedin, domaine = texte_js(valeurs.get("linkedinEntreprise")), texte_js(valeurs.get("siteEntreprise"))
    siren, nom = texte_js(valeurs.get("siren")), normaliser_nom_entreprise(valeurs.get("entreprise"))
    if linkedin and linkedin in index["linkedin"]:
        return "URL LinkedIn entreprise"
    if domaine and domaine in index["domaines"]:
        return "domaine"
    if siren and siren in index["sirens"]:
        return "SIREN"
    if nom and nom in index["noms"]:
        return "nom normalisé"
    return None


def _doublon(valeurs: dict, vus: dict, ligne: int) -> int | None:
    identifiants = [
        f"linkedin:{valeurs['linkedinProfil']}" if vrai(valeurs.get("linkedinProfil")) else "",
        f"email:{valeurs['emailProfessionnel']}" if vrai(valeurs.get("emailProfessionnel")) else "",
        f"mobile:{valeurs['mobile']}" if vrai(valeurs.get("mobile")) else "",
        f"contact:{normaliser_texte(valeurs['nomComplet'])}|{normaliser_texte(valeurs['entreprise'])}" if vrai(valeurs.get("nomComplet")) and vrai(valeurs.get("entreprise")) else "",
    ]
    identifiants = [i for i in identifiants if i]
    for identifiant in identifiants:
        if identifiant in vus:
            return vus[identifiant]
    for identifiant in identifiants:
        vus[identifiant] = ligne
    return None


def _trouver_attendu(attendus: list[str], texte: str) -> str | None:
    mots = set(texte.split(" "))
    for attendu in attendus:
        if (len(attendu) <= 3 and " " not in attendu and attendu in mots) or ((len(attendu) > 3 or " " in attendu) and attendu in texte):
            return attendu
    return None


def _evaluer_valeur(regle: dict, champ: str | None, normalisee: Any, blacklist: str | None, doublon: int | None) -> tuple[bool, str]:
    operateur = regle.get("operateur")
    texte = normaliser_texte(normalisee)
    if operateur == "pas_blackliste":
        return (not blacklist, f"Entreprise présente dans la blacklist ({blacklist})" if blacklist else "Entreprise absente de la blacklist")
    if operateur == "unique":
        return (not doublon, f"Doublon de la ligne {doublon}" if doublon else "Contact unique")
    if operateur == "present":
        rempli = rogner(texte_js(normalisee)) != ""
        return (rempli, "Valeur renseignée" if rempli else "Valeur absente")
    if operateur == "email_professionnel":
        ok = email_est_professionnel(normalisee)
        return (ok, "Email professionnel valide" if ok else "Email invalide ou personnel")
    if operateur == "mobile_francais":
        ok = bool(normaliser_mobile_francais(normalisee))
        return (ok, "Mobile français 06/07 valide" if ok else "Mobile français 06/07 invalide")
    if operateur == "provenance_professionnelle":
        perso = re.search(r"\b(perso|personnel|personnelle|personal|private)\b", texte) is not None
        pro = re.search(r"\b(pro|professionnel|professionnelle|work|business|company)\b", texte) is not None
        return (pro and not perso, "La provenance indique une donnée personnelle" if perso else "Provenance professionnelle confirmée" if pro else "Provenance professionnelle non démontrée")
    if operateur == "entre":
        intervalle = regle.get("valeur") if isinstance(regle.get("valeur"), dict) else {}
        borne_min, borne_max = intervalle.get("min"), intervalle.get("max")
        nombre = normalisee if est_nombre(normalisee) else lire_nombre(normalisee)
        try:
            ok = nombre is not None and float(nombre) >= float(borne_min) and float(nombre) <= float(borne_max)
        except (TypeError, ValueError):
            ok = False
        affiche = lambda v: nombre_js(v) if est_nombre(v) else texte_js(v) if v is not None else "undefined"  # noqa: E731
        lu = affiche(nombre) if nombre is not None else "valeur illisible"
        intervalle_texte = f"{affiche(borne_min)}–{affiche(borne_max)}"
        return (ok, f"{lu} est dans l'intervalle {intervalle_texte}" if ok else f"{lu} est hors de l'intervalle {intervalle_texte}")
    valeur = regle.get("valeur")
    attendus = [a for a in (normaliser_texte(v) for v in valeur) if a] if isinstance(valeur, list) else []
    if operateur == "contient_un":
        trouve = _trouver_attendu(attendus, texte)
        return (trouve is not None, f"Correspondance : {trouve}" if trouve is not None else "Aucun mot-clé accepté")
    if operateur == "ne_contient_aucun":
        trouve = _trouver_attendu(attendus, texte)
        return (trouve is None, f"Expression exclue : {trouve}" if trouve is not None else "Aucune expression exclue")
    assistant = champ == "poste" and re.search(r"\b(assistant|assistante|chief of staff|bras droit)\b", texte) is not None
    trouve = None
    if not assistant:
        for attendu in attendus:
            if (texte == attendu or texte.startswith(f"{attendu} ") or texte.endswith(f" {attendu}")) if champ == "poste" else texte == attendu:
                trouve = attendu
                break
    return (trouve is not None, f"Valeur acceptée : {trouve}" if trouve is not None else "Poste assistant exclu" if assistant else "Valeur hors liste")


def _evaluation(regle_id, critere_id, libelle, portee, importance, statut, brute, normalisee, colonne, raison) -> dict:
    return {
        "regleId": regle_id, "critereId": critere_id, "libelle": libelle, "portee": portee, "importance": importance,
        "statut": statut, "valeurBrute": brute, "valeurNormalisee": normalisee, "colonneSource": colonne, "raison": raison,
    }


def _evaluer_regle(regle: dict, source: dict, normalisees: dict, mapping: dict, blacklist: str | None, doublon: int | None) -> dict:
    critere_id = texte_js(regle.get("critereId"))
    definition = DEFINITIONS.get(critere_id)
    champ = definition[0] if definition else None
    colonne = (mapping.get(champ) or "") if champ else "Contrôle calculé"
    if champ and colonne:
        brute = source.get(colonne, INDEFINI)
    elif regle.get("operateur") == "pas_blackliste":
        brute = blacklist if blacklist is not None else "absent"
    elif regle.get("operateur") == "unique":
        brute = doublon if doublon is not None else "unique"
    else:
        brute = ""
    normalisee = normalisees.get(champ, "") if champ else texte_js(brute)
    importance = regle.get("importance")
    if not definition:
        return _evaluation(regle.get("id"), critere_id, critere_id, "qualite", importance, "Non évalué", texte_js(brute), texte_js(normalisee), colonne, "Critère inconnu")
    _, libelle, portee = definition
    if champ is not None and normalisee == "" and not est_nombre(normalisee):
        return _evaluation(regle.get("id"), critere_id, libelle, portee, importance, "Donnée absente", texte_js(brute), "", colonne, f"Donnée absente — politique : {texte_js(regle.get('siAbsent'))}")
    reussi, raison = _evaluer_valeur(regle, champ, normalisee, blacklist, doublon)
    return _evaluation(regle.get("id"), critere_id, libelle, portee, importance, "Réussi" if reussi else "Échoué", texte_js(brute), texte_js(normalisee), colonne, raison)


def valider_export_b2b(candidats: list[dict], blacklist: list[dict], configuration: dict) -> list[dict]:
    """Rend une ligne validée par candidat : statut, raisons, cause principale, évaluations, valeurs."""
    index = _index_blacklist(blacklist, configuration["mappingBlacklist"])
    vus: dict[str, int] = {}
    mapping = configuration["mappingCandidats"]
    politique = {r.get("id"): r.get("siAbsent") for r in reversed(configuration["regles"])}
    regles_actives = [r for r in configuration["regles"] if vrai(r.get("active"))]
    lignes = []
    for position, source in enumerate(candidats):
        numero = position + 2
        valeurs = valeurs_normalisees(source, mapping)
        trouvee = _correspondance_blacklist(valeurs, index)
        doublon = _doublon(valeurs, vus, numero)
        evaluations = []
        for champ in configuration["champsObligatoires"]:
            colonne = mapping.get(champ) or ""
            brute = source.get(colonne, INDEFINI) if colonne else ""
            presente = champ in valeurs
            evaluations.append(_evaluation(
                f"obligatoire:{champ}", f"obligatoire:{champ}", f"Champ obligatoire : {champ}",
                "contact" if champ in CHAMPS_CONTACT else "entreprise", "bloquant",
                "Réussi" if presente else "Donnée absente", texte_js(brute), texte_js(valeurs.get(champ)), colonne,
                "Champ renseigné" if presente else "Champ obligatoire absent",
            ))
        evaluations += [_evaluer_regle(r, source, valeurs, mapping, trouvee, doublon) for r in regles_actives]

        def politique_de(evaluation: dict, defaut: str) -> str:
            valeur = politique.get(evaluation["regleId"], INDEFINI)
            return defaut if valeur is None or valeur is INDEFINI else valeur

        rejette = lambda e: e["importance"] == "bloquant" and (e["statut"] in ("Échoué", "Non évalué") or (e["statut"] == "Donnée absente" and politique_de(e, "rejeter") == "rejeter"))  # noqa: E731
        fait_verifier = lambda e: e["statut"] == "Non évalué" or (e["statut"] == "Échoué" and e["importance"] == "verification") or (e["statut"] == "Donnée absente" and politique_de(e, "verifier") == "verifier")  # noqa: E731
        cause_rejet = next((e for e in evaluations if rejette(e)), None)
        cause_verification = None if cause_rejet else next((e for e in evaluations if fait_verifier(e)), None)
        statut = "Rejeté" if cause_rejet else "À vérifier" if cause_verification else "Conforme"
        lignes.append({
            "numeroLigne": numero, "statut": statut,
            "raisons": [f"{e['libelle']} : {e['raison']}" for e in evaluations if e["statut"] != "Réussi"],
            "causePrincipale": cause_rejet or cause_verification,
            "evaluations": evaluations, "valeurs": valeurs, "source": source,
        })
    return lignes


def type_probleme(ligne: dict) -> str:
    """Libellé de la cause principale, comme la colonne « Type de problème » de la page."""
    evaluation = ligne.get("causePrincipale")
    if not evaluation:
        return ""
    critere = evaluation["critereId"]
    if critere == "blacklist":
        return "Blacklist"
    if critere == "doublon":
        return "Doublon"
    if critere in ("poste", "obligatoire:poste"):
        return "Mauvais intitulé de poste"
    if critere in ("activite", "tags_source", "qualification_activite"):
        return "Entreprise à vérifier" if ligne["statut"] == "À vérifier" else "Mauvaise entreprise"
    if critere.startswith("obligatoire:"):
        return f"Champ obligatoire manquant — {re.sub(r'^Champ obligatoire[ ]*:[ ]*', '', evaluation['libelle'], flags=re.I)}"
    return evaluation["libelle"]


def bilan_b2b(lignes: list[dict]) -> dict:
    """Le bilan obligatoire après chaque passe : totaux, une cause principale par ligne non conforme."""
    comptes = Counter(ligne["statut"] for ligne in lignes)
    causes: dict[str, Counter] = {"À vérifier": Counter(), "Rejeté": Counter()}
    for ligne in lignes:
        if ligne["statut"] != "Conforme":
            causes[ligne["statut"]][type_probleme(ligne)] += 1
    conformes, a_verifier, rejetes = comptes["Conforme"], comptes["À vérifier"], comptes["Rejeté"]
    somme_causes = sum(sum(c.values()) for c in causes.values())
    return {
        "total": len(lignes), "conformes": conformes, "a_verifier": a_verifier, "rejetes": rejetes,
        "causes": {statut: dict(compteur.most_common()) for statut, compteur in causes.items()},
        "invariants": {
            "total = conformes + à vérifier + rejetés": len(lignes) == conformes + a_verifier + rejetes,
            "à vérifier + rejetés = somme des causes": a_verifier + rejetes == somme_causes,
        },
    }


# --------------------------------------------------------------------------- personas (jumeau de engine.ts)

NON_CLASSE = "Non catégorisé"
COLONNES_POSTE = ("poste", "intitule de poste", "intitule du poste", "intitule", "fonction", "job title", "title", "position")


def normaliser_intitule(texte: Any) -> str:
    if texte is None or texte is INDEFINI:
        return ""
    return rogner(sans_diacritiques(texte_js(texte)).lower())


def _mot_entier(cible: str, cle: str) -> bool:
    return re.search(r"(?<![a-z0-9])" + re.escape(cle) + r"(?![a-z0-9])", cible) is not None


def classer(intitule: Any, axe: dict, multi: bool = False) -> str:
    cible = normaliser_intitule(intitule)
    if not cible:
        return NON_CLASSE
    trouvees = []
    for regle in axe.get("regles", []):
        for mot in regle.get("motsCles", []):
            cle = normaliser_intitule(mot)
            if cle and (_mot_entier(cible, cle) if len(cle) <= 3 else cle in cible):
                trouvees.append((_defaut(regle.get("priorite"), 999), regle["categorie"]))
                break
    if not trouvees:
        return NON_CLASSE
    trouvees.sort(key=lambda t: t[0])
    if not multi:
        return trouvees[0][1]
    return ", ".join(OrderedDict.fromkeys(categorie for _, categorie in trouvees))


def segmenter(intitules: list, matrice: dict, multi: bool = False) -> dict:
    axes = matrice.get("axes", [])
    lignes, distributions, manquants = [], {a["id"]: Counter() for a in axes}, {a["id"]: Counter() for a in axes}
    for intitule in intitules:
        valeurs = {}
        for axe in axes:
            categorie = classer(intitule, axe, multi)
            valeurs[axe["id"]] = categorie
            distributions[axe["id"]][categorie] += 1
            if categorie == NON_CLASSE and normaliser_intitule(intitule):
                manquants[axe["id"]][texte_js(intitule)] += 1
        for derive in matrice.get("derives", []) or []:
            valeurs[derive["id"]] = derive.get("table", {}).get(valeurs.get(derive["depuis"], ""), _defaut(derive.get("defaut"), ""))
        lignes.append(valeurs)
    return {"lignes": lignes, "distributions": distributions, "nonClasses": manquants, "total": len(intitules)}


def deviner_colonne_poste(entetes: list[str]) -> int:
    normalisees = [normaliser_intitule(e) for e in entetes]
    for connue in COLONNES_POSTE:
        if connue in normalisees:
            return normalisees.index(connue)
    return next((i for i, e in enumerate(normalisees) if "poste" in e or "titre" in e or "title" in e), -1)


# --------------------------------------------------------------------------- fichiers


class ErreurEntree(Exception):
    """Une entrée illisible : on s'arrête avec un message clair, jamais avec une trace."""


def _cellule(valeur: Any) -> Any:
    if valeur is None:
        return ""
    if isinstance(valeur, (dt.datetime, dt.date)):
        return valeur.isoformat()
    return valeur


def lire_tableau(chemin: Path) -> tuple[list[str], list[list[Any]]]:
    """(entêtes, lignes) d'un .xlsx, .xlsm ou .csv ; les lignes vides sont ignorées, comme dans la page."""
    if not chemin.exists():
        raise ErreurEntree(f"Fichier introuvable : {chemin}")
    if chemin.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
        except ImportError as erreur:  # pragma: no cover - dépend de l'environnement
            raise ErreurEntree("openpyxl est requis pour lire un fichier Excel : python -m pip install openpyxl") from erreur
        classeur = openpyxl.load_workbook(chemin, data_only=True, read_only=True)
        brutes = [[_cellule(c) for c in ligne] for ligne in classeur.worksheets[0].iter_rows(values_only=True)]
        classeur.close()
    elif chemin.suffix.lower() in {".csv", ".txt", ".tsv"}:
        octets = chemin.read_bytes()
        try:
            texte = octets.decode("utf-8-sig")
        except UnicodeDecodeError:
            texte = octets.decode("cp1252", errors="replace")
        premiere = texte.split("\n", 1)[0]
        separateur = None
        if premiere.lower().startswith("sep=") and len(premiere.strip()) >= 5:
            separateur, texte = premiere.strip()[4], texte.split("\n", 1)[1] if "\n" in texte else ""
        if separateur is None:
            separateur = max(",;\t|", key=lambda s: premiere.count(s)) if any(s in premiere for s in ",;\t|") else ","
        brutes = [list(ligne) for ligne in csv.reader(io.StringIO(texte, newline=""), delimiter=separateur)]
    else:
        raise ErreurEntree(f"Format non pris en charge : {chemin.name} (attendu .xlsx, .xlsm ou .csv)")
    lignes = [ligne for ligne in brutes if any(rogner(texte_js(c)) for c in ligne)]
    if not lignes:
        return [], []
    entetes = [texte_js(c) for c in lignes[0]]
    return entetes, lignes[1:]


def vers_objets(entetes: list[str], lignes: list[list[Any]]) -> list[dict]:
    """Comme la page : une ligne devient {entête: valeur} ; à entête répétée, la dernière colonne l'emporte."""
    return [{entete: (ligne[i] if i < len(ligne) else "") for i, entete in enumerate(entetes)} for ligne in lignes]


def securiser(valeur: Any) -> Any:
    """Neutralise une cellule qui commencerait comme une formule de tableur."""
    if isinstance(valeur, str) and re.match(r"^[=+\-@]", valeur):
        return f"'{valeur}"
    return valeur


def _openpyxl():
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as erreur:  # pragma: no cover
        raise ErreurEntree("openpyxl est requis pour écrire le classeur : python -m pip install openpyxl") from erreur
    return openpyxl, Alignment, Font, PatternFill, get_column_letter, DataValidation, Table, TableStyleInfo


def _feuille(classeur, titre: str, nom_table: str, entetes: list[str], lignes: list[list[Any]], largeurs: list[int]):
    _, Alignment, Font, PatternFill, lettre, _, Table, TableStyleInfo = _openpyxl()
    feuille = classeur.create_sheet(titre)
    feuille.append(list(entetes))
    for ligne in lignes:
        feuille.append([c[0] if isinstance(c, tuple) else c for c in ligne])
    for numero, ligne in enumerate(lignes, start=2):
        for colonne, cellule in enumerate(ligne, start=1):
            if isinstance(cellule, tuple):
                feuille.cell(numero, colonne).hyperlink = cellule[1]
                feuille.cell(numero, colonne).style = "Hyperlink"
    if lignes:
        table = Table(displayName=nom_table, ref=f"A1:{lettre(len(entetes))}{len(lignes) + 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        feuille.add_table(table)
    feuille.freeze_panes = "A2"
    for index, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[lettre(index)].width = largeur
    for cellule in feuille[1]:
        cellule.font = Font(bold=True, color="FFFFFFFF")
        cellule.fill = PatternFill("solid", fgColor="FF164E63")
        cellule.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    feuille.row_dimensions[1].height = 30
    return feuille


ENTETES_CONTROLE_QA = [
    "Verdict", "Nom complet", "Entreprise", "Poste", "LinkedIn profil", "Catégorie LinkedIn", "Employés visibles",
    "Email professionnel", "Mobile 06/07", "Type de problème", "Pourquoi", "Source entreprise", "Source LinkedIn",
    "Contrôlé le", "N° ligne source",
]


def _lien(valeur: Any) -> Any:
    texte = rogner(texte_js(valeur))
    return (texte, texte) if re.match(r"^https?://", texte, re.I) else texte_js(securiser(texte))


def ecrire_classeur_b2b(chemin: Path, lignes: list[dict], configuration: dict, prevol: dict, mention: str) -> None:
    openpyxl, Alignment, Font, PatternFill, _, DataValidation, _, _ = _openpyxl()
    classeur = openpyxl.Workbook()
    classeur.remove(classeur.active)
    mapping = configuration["mappingCandidats"]
    brute = lambda ligne, champ: ligne["source"].get(mapping[champ], "") if mapping.get(champ) else ligne["valeurs"].get(champ, "")  # noqa: E731
    sur = lambda v: v if est_nombre(v) else texte_js(securiser(texte_js(v)))  # noqa: E731
    controle = [[
        l["statut"], sur(l["valeurs"].get("nomComplet", "")), sur(l["valeurs"].get("entreprise", "")), sur(l["valeurs"].get("poste", "")),
        _lien(l["valeurs"].get("linkedinProfil", "")), sur(l["valeurs"].get("categorieLinkedIn", "")), sur(l["valeurs"].get("employesVisibles", "")),
        sur(l["valeurs"].get("emailProfessionnel", "")), sur(l["valeurs"].get("mobile", "")), type_probleme(l),
        " | ".join(texte_js(securiser(r)) for r in l["raisons"]), _lien(brute(l, "siteEntreprise")), _lien(brute(l, "linkedinEntreprise")),
        None, l["numeroLigne"],
    ] for l in lignes]
    feuille = _feuille(classeur, "Contrôle QA", "ControleQA", ENTETES_CONTROLE_QA, controle, [16, 23, 25, 24, 36, 18, 18, 30, 18, 29, 56, 34, 38, 16, 16])
    validation = DataValidation(type="list", formula1='"Conforme,À vérifier,Rejeté"', allow_blank=False, showErrorMessage=True, errorTitle="Verdict invalide", error="Choisir Conforme, À vérifier ou Rejeté.")
    feuille.add_data_validation(validation)
    couleurs = {"Conforme": ("FFDCFCE7", "FF166534"), "Rejeté": ("FFFEE2E2", "FF991B1B")}
    for numero in range(2, len(lignes) + 2):
        cellule = feuille.cell(numero, 1)
        fond, texte = couleurs.get(cellule.value, ("FFFEF3C7", "FF92400E"))
        cellule.fill = PatternFill("solid", fgColor=fond)
        cellule.font = Font(bold=True, color=texte)
        cellule.alignment = Alignment(vertical="center", horizontal="center")
        validation.add(cellule)
        feuille.cell(numero, 14).number_format = "yyyy-mm-dd"
        for colonne in (11, 12, 13):
            feuille.cell(numero, colonne).alignment = Alignment(vertical="top", wrap_text=True)

    audit = [[
        l["numeroLigne"], l["statut"], e["portee"], e["libelle"], e["importance"], e["statut"], e["colonneSource"],
        sur(e["valeurBrute"]), sur(e["valeurNormalisee"]), sur(e["raison"]),
    ] for l in lignes for e in l["evaluations"]]
    _feuille(classeur, "Audit moulinette", "AuditMoulinette",
             ["N° ligne source", "Statut moulinette", "Portée", "Critère", "Importance", "Résultat critère", "Colonne source", "Valeur brute", "Valeur normalisée", "Diagnostic"],
             audit, [16, 20, 16, 32, 15, 20, 24, 30, 30, 48])

    bilan = bilan_b2b(lignes)
    resume = [
        ["Version du schéma", configuration.get("schemaVersion", 2)],
        ["Lignes analysées", bilan["total"]],
        ["Conformes moulinette", bilan["conformes"]],
        ["À vérifier", bilan["a_verifier"]],
        ["Rejetées", bilan["rejetes"]],
        ["Contrôle effectué le", dt.datetime.now().replace(microsecond=0)],
        ["Empreinte blacklist", texte_js(configuration["blacklist"].get("empreinte")) or "Non mémorisée"],
        *[[f"Critère — {DEFINITIONS.get(r.get('critereId'), (None, r.get('critereId')))[1]}",
           f"{r.get('operateur')} | {json.dumps(r.get('valeur'), ensure_ascii=False, separators=(',', ':'))} | {r.get('importance')} | absent={r.get('siAbsent')}"]
          for r in configuration["regles"] if vrai(r.get("active"))],
        ["Prévol", "autorisé" if prevol["autorise"] else f"BLOQUÉ — {', '.join(c['id'] for c in prevol['controles'] if c['bloquant'] and not c['valide'])}"],
        ["Nature du passage", mention],
        ["Outil", f"moulinette.py {VERSION}"],
    ]
    feuille = _feuille(classeur, "Résumé", "ResumeMoulinette", ["Indicateur", "Valeur"], resume, [46, 90])
    feuille["B7"].number_format = "yyyy-mm-dd hh:mm"
    classeur.save(chemin)


def ecrire_conformes_b2b(chemin: Path, lignes: list[dict], configuration: dict) -> None:
    openpyxl = _openpyxl()[0]
    classeur = openpyxl.Workbook()
    classeur.remove(classeur.active)
    colonnes = configuration["colonnesFinales"]
    entetes = [c.get("libelle", c.get("champ")) for c in colonnes]
    valeurs = [[(v if est_nombre(v) else texte_js(securiser(texte_js(v)))) for v in (l["valeurs"].get(c.get("champ"), "") for c in colonnes)] for l in lignes if l["statut"] == "Conforme"]
    _feuille(classeur, "Conformes", "ConformesClient", entetes, valeurs, [34 if re.search(r"linkedin|email", e, re.I) else 22 for e in entetes])
    classeur.save(chemin)


def ecrire_personas(chemin: Path, entetes: list[str], lignes: list[list[Any]]) -> None:
    openpyxl = _openpyxl()[0]
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.title = "Données qualifiées"
    feuille.append(entetes)
    for ligne in lignes:
        feuille.append([v if est_nombre(v) else texte_js(securiser(texte_js(v))) for v in ligne])
    classeur.save(chemin)


# --------------------------------------------------------------------------- commandes


def _base(chemin: Path) -> str:
    return re.sub(r"\.(xlsx|xlsm|xls|csv)$", "", chemin.name, flags=re.I) or "export"


def _ecrire_sortie(donnees: dict, en_json: bool, texte: list[str]) -> None:
    if en_json:
        sys.stdout.write(json.dumps(donnees, ensure_ascii=False, indent=2) + "\n")
    else:
        sys.stdout.write("\n".join(texte) + "\n")


def _charger_b2b(args) -> tuple[dict, dict, list[dict], list[dict], list[str]]:
    identite, configuration = charger_mission(args.moulinette)
    if configuration["mode"] != "revalidation_b2b":
        raise ErreurEntree("Cette mission est une matrice de personas : utilise la commande « personas ».")
    entetes, lignes = lire_tableau(args.export)
    if not entetes:
        raise ErreurEntree(f"{args.export.name} est vide.")
    configuration["mappingCandidats"] = completer_mapping(configuration["mappingCandidats"], entetes, ENTETES_CANDIDATS)
    blacklist: list[dict] = []
    if args.blacklist:
        entetes_bl, lignes_bl = lire_tableau(args.blacklist)
        configuration["mappingBlacklist"] = completer_mapping(configuration["mappingBlacklist"], entetes_bl, ENTETES_BLACKLIST)
        blacklist = vers_objets(entetes_bl, lignes_bl)
    return identite, configuration, vers_objets(entetes, lignes), blacklist, entetes


def _prevol(configuration: dict, candidats: list, blacklist: list) -> dict:
    empreinte = empreinte_blacklist(blacklist, configuration["mappingBlacklist"]) if blacklist else None
    controles = controler_prevol(configuration, len(candidats), len(blacklist), empreinte)
    return {"autorise": prevol_autorise(controles), "controles": controles, "empreinte_blacklist_chargee": empreinte}


def _nom_mission(identite: dict, chemin: Path) -> str:
    return " / ".join(v for v in (identite.get("client"), identite.get("mission")) if v) or chemin.stem


def commande_b2b(args) -> int:
    identite, configuration, candidats, blacklist, _ = _charger_b2b(args)
    prevol = _prevol(configuration, candidats, blacklist)
    nom = _nom_mission(identite, args.moulinette)
    bloques = [c for c in prevol["controles"] if c["bloquant"] and not c["valide"]]
    if not prevol["autorise"] and not args.diagnostic:
        _ecrire_sortie(
            {"outil": "moulinette", "version": VERSION, "commande": "b2b", "mission": identite, "prevol": prevol, "livre": False},
            args.json,
            [f"Moulinette « {nom} » — {args.export.name}", f"Prévol BLOQUÉ : {len(bloques)} contrôle(s) à lever, rien n'est livré."]
            + [f"  ✗ {c['libelle']}" for c in bloques]
            + ["Pour voir les verdicts sans rien livrer : ajoute --diagnostic."],
        )
        return 2
    lignes = valider_export_b2b(candidats, blacklist, configuration)
    bilan = bilan_b2b(lignes)
    dossier = args.sortie or args.export.parent
    dossier.mkdir(parents=True, exist_ok=True)
    diagnostic = not prevol["autorise"]
    suffixe = "-diagnostic" if diagnostic else ""
    fichiers = {"controle_qa": str(dossier / f"{_base(args.export)}{suffixe}-controle-qa.xlsx")}
    mention = "Diagnostic non livrable (prévol non passé)" if diagnostic else "Passage livrable (prévol passé)"
    ecrire_classeur_b2b(Path(fichiers["controle_qa"]), lignes, configuration, prevol, mention)
    if not diagnostic:
        fichiers["conformes"] = str(dossier / f"{_base(args.export)}-conformes.xlsx")
        ecrire_conformes_b2b(Path(fichiers["conformes"]), lignes, configuration)
    sortie = {
        "outil": "moulinette", "version": VERSION, "commande": "b2b", "mission": identite,
        "prevol": prevol, "diagnostic": diagnostic, "livre": not diagnostic, **bilan, "fichiers": fichiers,
    }
    if args.lignes:
        sortie["lignes"] = [{"numeroLigne": l["numeroLigne"], "statut": l["statut"], "cause": type_probleme(l), "raisons": l["raisons"]} for l in lignes]
    causes = lambda statut: " · ".join(f"{cause} {n}" for cause, n in bilan["causes"][statut].items()) or "aucune"  # noqa: E731
    _ecrire_sortie(sortie, args.json, [
        f"Moulinette « {nom} » — {args.export.name}",
        "Prévol : autorisé" if not diagnostic else f"Prévol non passé ({len(bloques)} contrôle(s)) : DIAGNOSTIC NON LIVRABLE",
        f"Lignes {bilan['total']} · conformes {bilan['conformes']} · à vérifier {bilan['a_verifier']} · rejetées {bilan['rejetes']}",
        f"Rejetées, par cause : {causes('Rejeté')}",
        f"À vérifier, par cause : {causes('À vérifier')}",
        f"Contrôles du bilan : {'justes' if all(bilan['invariants'].values()) else 'FAUX'}",
        "Écrit : " + " · ".join(fichiers.values()),
    ])
    return 0 if all(bilan["invariants"].values()) else 1


def commande_prevol(args) -> int:
    identite, configuration, candidats, blacklist, _ = _charger_b2b(args)
    prevol = _prevol(configuration, candidats, blacklist)
    _ecrire_sortie(
        {"outil": "moulinette", "version": VERSION, "commande": "prevol", "mission": identite, "prevol": prevol},
        args.json,
        [f"Prévol « {_nom_mission(identite, args.moulinette)} » : {'autorisé' if prevol['autorise'] else 'BLOQUÉ'}"]
        + [f"  {'✓' if c['valide'] else '✗' if c['bloquant'] else '·'} {c['libelle']}" for c in prevol["controles"]],
    )
    return 0 if prevol["autorise"] else 2


def commande_empreintes(args) -> int:
    identite, configuration = charger_mission(args.moulinette)
    if configuration["mode"] != "revalidation_b2b":
        raise ErreurEntree("Les empreintes ne concernent qu'une mission B2B.")
    empreintes = {
        "strategie": empreinte_strategie(configuration["strategie"]),
        "recherche": empreinte_recherche(configuration),
        "filtres": empreinte_filtres(configuration["pharow"]["filtres"]),
        "erreurs_strategie": erreurs_strategie(configuration["strategie"]),
    }
    if args.blacklist:
        entetes, lignes = lire_tableau(args.blacklist)
        mapping = completer_mapping(configuration["mappingBlacklist"], entetes, ENTETES_BLACKLIST)
        objets = vers_objets(entetes, lignes)
        empreintes.update({"blacklist": empreinte_blacklist(objets, mapping), "blacklist_volume": len(objets)})
    _ecrire_sortie({"outil": "moulinette", "version": VERSION, "commande": "empreintes", "mission": identite, **empreintes}, args.json,
                   [f"{cle} : {valeur}" for cle, valeur in empreintes.items()])
    return 0


def commande_personas(args) -> int:
    if args.matrice:
        _, configuration = charger_mission(args.matrice)
        if configuration["mode"] != "persona":
            raise ErreurEntree("Ce fichier est une mission B2B : utilise la commande « b2b ».")
        matrice, multi = configuration["matrice"], args.multi or configuration["multiCategories"]
    else:
        matrice, multi = json.loads(MATRICE_DEFAUT.read_text(encoding="utf-8")), args.multi
    entetes, lignes = lire_tableau(args.entree)
    if not entetes:
        raise ErreurEntree(f"{args.entree.name} est vide.")
    if args.colonne:
        cibles = [i for i, e in enumerate(entetes) if normaliser_intitule(e) == normaliser_intitule(args.colonne)]
        if not cibles:
            raise ErreurEntree(f"Colonne introuvable : {args.colonne!r}. Entêtes lues : {entetes}")
        index = cibles[0]
    else:
        index = deviner_colonne_poste(entetes)
        if index < 0:
            raise ErreurEntree(f"Impossible de deviner la colonne des intitulés de poste : précise-la avec -c. Entêtes lues : {entetes}")
    intitules = [ligne[index] if index < len(ligne) else "" for ligne in lignes]
    resultat = segmenter(intitules, matrice, multi)
    cles = [a["id"] for a in matrice.get("axes", [])] + [d["id"] for d in matrice.get("derives", []) or []]
    libelles = [a["libelle"] for a in matrice.get("axes", [])] + [d["libelle"] for d in matrice.get("derives", []) or []]
    destination = args.sortie or args.entree.with_name(f"{_base(args.entree)}-qualifie.xlsx")
    destination.parent.mkdir(parents=True, exist_ok=True)
    ecrire_personas(destination, list(entetes) + libelles, [list(l) + [v.get(c, "") for c in cles] for l, v in zip(lignes, resultat["lignes"])])
    axes = []
    for axe in matrice.get("axes", []):
        distribution = resultat["distributions"][axe["id"]]
        axes.append({
            "id": axe["id"], "libelle": axe["libelle"], "classees": resultat["total"] - distribution.get(NON_CLASSE, 0),
            "distribution": dict(distribution.most_common()),
            "non_classes": dict(resultat["nonClasses"][axe["id"]].most_common(20)),
        })
    texte = [f"Colonne lue : {entetes[index]!r} — {resultat['total']} lignes"]
    for axe in axes:
        detail = " | ".join(f"{k} {v}" for k, v in axe["distribution"].items() if k != NON_CLASSE)
        texte.append(f"  {axe['libelle']} : {axe['classees']}/{resultat['total']} classées — {detail}")
        if axe["non_classes"]:
            texte.append("    non classés : " + " | ".join(f"{k} ({v})" for k, v in list(axe["non_classes"].items())[:5]))
    texte.append(f"Écrit : {destination}")
    _ecrire_sortie({"outil": "moulinette", "version": VERSION, "commande": "personas", "colonne": entetes[index], "total": resultat["total"], "axes": axes, "fichier": str(destination)}, args.json, texte)
    return 0


def construire_parseur() -> argparse.ArgumentParser:
    parseur = argparse.ArgumentParser(prog="moulinette", description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parseur.add_argument("--version", action="version", version=f"moulinette {VERSION}")
    sous = parseur.add_subparsers(dest="commande", required=True)

    personas = sous.add_parser("personas", help="trier une colonne d'intitulés de poste")
    personas.add_argument("entree", type=Path, help="fichier .xlsx, .xlsm ou .csv")
    personas.add_argument("-m", "--matrice", type=Path, help="matrice JSON ou mission « persona » (défaut : matrice générique)")
    personas.add_argument("-o", "--sortie", type=Path, help="fichier de sortie (défaut : <entrée>-qualifie.xlsx)")
    personas.add_argument("-c", "--colonne", help="nom exact de la colonne des intitulés (sinon devinée)")
    personas.add_argument("--multi", action="store_true", help="garder toutes les catégories trouvées, par priorité")
    personas.add_argument("--json", action="store_true", help="bilan JSON sur la sortie standard (pour un agent)")
    personas.set_defaults(action=commande_personas)

    for nom, aide, action in (
        ("b2b", "revalider un export selon la moulinette d'une mission", commande_b2b),
        ("prevol", "dire seulement si la mission peut tourner", commande_prevol),
    ):
        commande = sous.add_parser(nom, help=aide)
        commande.add_argument("export", type=Path, help="export candidats .xlsx, .xlsm ou .csv")
        commande.add_argument("--moulinette", type=Path, required=True, help="fichier JSON de la mission")
        commande.add_argument("--blacklist", type=Path, help="liste d'exclusion .xlsx ou .csv")
        commande.add_argument("--json", action="store_true", help="bilan JSON sur la sortie standard (pour un agent)")
        if nom == "b2b":
            commande.add_argument("-o", "--sortie", type=Path, help="dossier de sortie (défaut : celui de l'export)")
            commande.add_argument("--diagnostic", action="store_true", help="rendre les verdicts même si le prévol bloque ; rien de livrable")
            commande.add_argument("--lignes", action="store_true", help="ajouter le verdict de chaque ligne au bilan JSON")
        commande.set_defaults(action=action)

    empreintes = sous.add_parser("empreintes", help="calculer les empreintes que la mission mémorise")
    empreintes.add_argument("--moulinette", type=Path, required=True, help="fichier JSON de la mission")
    empreintes.add_argument("--blacklist", type=Path, help="liste d'exclusion, pour son empreinte et son volume")
    empreintes.add_argument("--json", action="store_true", help="sortie JSON")
    empreintes.set_defaults(action=commande_empreintes)
    return parseur


def main(argv: list[str] | None = None) -> int:
    for flux in (sys.stdout, sys.stderr):
        try:
            flux.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):  # pragma: no cover
            pass
    args = construire_parseur().parse_args(argv)
    try:
        return args.action(args)
    except ErreurEntree as erreur:
        sys.stderr.write(f"moulinette : {erreur}\n")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
