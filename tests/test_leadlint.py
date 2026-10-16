# -*- coding: utf-8 -*-
"""Banc de leadlint : parité avec le moteur de la page web, puis la ligne de commande de bout en bout.

    python tests/test_leadlint.py

Les attendus de `tests/parite/` sont rendus par le moteur de la page web (oracle
`scripts/segmentation/oracle_parite.ts` du dépôt de la Marketplace). Le banc rejoue les
mêmes entrées dans le script et exige le même résultat, champ par champ. Le pont de bout
en bout : la mission que la page fait télécharger (`mission-telechargee.json`) passe le
prévol du script et en rend les empreintes. Des contrôles sont négatifs : un banc
incapable de dire non ne vaut pas son oui.
"""

from __future__ import annotations

import copy
import contextlib
import csv
import io
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

RACINE = Path(__file__).resolve().parents[1]
PARITE = RACINE / "tests" / "parite"
EXEMPLES = RACINE / "skills" / "leadlint" / "exemples"
SCRIPT = RACINE / "skills" / "leadlint" / "scripts" / "leadlint.py"
sys.path.insert(0, str(SCRIPT.parent))

import leadlint as m  # noqa: E402


def lire(nom: str):
    return json.loads((PARITE / nom).read_text(encoding="utf-8"))


def rendu_b2b(entree: dict) -> dict:
    """Ce que le script rend pour une entrée de parité, au format de l'oracle."""
    configuration = m.migrer_configuration_b2b(copy.deepcopy(entree["configuration"]))
    candidats, blacklist = entree["candidats"], entree["blacklist"]
    empreinte_bl = m.empreinte_blacklist(blacklist, configuration["mappingBlacklist"]) if blacklist else None
    controles = m.controler_prevol(configuration, len(candidats), len(blacklist), empreinte_bl)
    lignes = m.valider_export_b2b(candidats, blacklist, configuration)
    return {
        "empreintes": {
            "strategie": m.empreinte_strategie(configuration["strategie"]), "recherche": m.empreinte_recherche(configuration),
            "filtres": m.empreinte_filtres(configuration["pharow"]["filtres"]), "blacklist": empreinte_bl,
        },
        "prevol": {"autorise": m.prevol_autorise(controles), "controles": controles},
        "lignes": [{
            "numeroLigne": l["numeroLigne"], "statut": l["statut"], "cause": l["causePrincipale"]["critereId"] if l["causePrincipale"] else None,
            "typeProbleme": m.type_probleme(l), "raisons": l["raisons"], "valeurs": l["valeurs"], "evaluations": l["evaluations"],
        } for l in lignes],
    }


def ecarts(attendu, obtenu, chemin="") -> list[str]:
    """Toutes les différences, avec leur chemin ; un nombre vaut l'autre s'il est égal (20 == 20.0)."""
    if isinstance(attendu, dict) and isinstance(obtenu, dict):
        sortie = []
        for cle in sorted(set(attendu) | set(obtenu)):
            if cle not in obtenu or cle not in attendu:
                sortie.append(f"{chemin}.{cle} : présent d'un seul côté")
            else:
                sortie += ecarts(attendu[cle], obtenu[cle], f"{chemin}.{cle}")
        return sortie
    if isinstance(attendu, list) and isinstance(obtenu, list):
        if len(attendu) != len(obtenu):
            return [f"{chemin} : {len(attendu)} éléments attendus, {len(obtenu)} obtenus"]
        return [e for i, (a, o) in enumerate(zip(attendu, obtenu)) for e in ecarts(a, o, f"{chemin}[{i}]")]
    if m.est_nombre(attendu) and m.est_nombre(obtenu):
        return [] if attendu == obtenu else [f"{chemin} : {attendu!r} attendu, {obtenu!r} obtenu"]
    return [] if attendu == obtenu and type(attendu) is type(obtenu) else [f"{chemin} : {attendu!r} attendu, {obtenu!r} obtenu"]


SCENARIOS_B2B = ["esn-validee", "esn-non-validee", "esn-filtre-modifie", "regles-variees"]


class Parite(unittest.TestCase):
    """Le script rend exactement ce que rend la page web."""

    def test_scenarios_b2b(self):
        for nom in SCENARIOS_B2B:
            with self.subTest(scenario=nom):
                differences = ecarts(lire(f"b2b-{nom}.attendu.json"), rendu_b2b(lire(f"b2b-{nom}.entree.json")))
                self.assertEqual(differences, [], "\n".join(differences[:15]))

    def test_mission_validee_passe_le_prevol_apres_rechargement(self):
        self.assertTrue(lire("b2b-esn-validee.attendu.json")["prevol"]["autorise"])
        self.assertTrue(rendu_b2b(lire("b2b-esn-validee.entree.json"))["prevol"]["autorise"])

    def test_ancienne_mission(self):
        entree, attendu = lire("b2b-ancienne.entree.json"), lire("b2b-ancienne.attendu.json")
        configuration = m.migrer_configuration_b2b(entree["configuration"])
        regles = [{k: r[k] for k in ("critereId", "operateur", "valeur", "importance", "siAbsent")} for r in configuration["regles"]]
        self.assertEqual(regles, attendu["regles"])
        self.assertEqual([l["statut"] for l in m.valider_export_b2b(entree["candidats"], [], configuration)], attendu["statuts"])

    def test_normaliseurs(self):
        fonctions = {
            "texte": m.normaliser_texte, "nomEntreprise": m.normaliser_nom_entreprise, "linkedin": m.normaliser_url_linkedin,
            "domaine": m.normaliser_domaine, "siren": m.normaliser_siren, "categorie": m.normaliser_categorie_linkedin,
            "nombre": m.lire_nombre, "email": m.normaliser_email, "emailPro": m.email_est_professionnel,
            "mobile": m.normaliser_mobile_francais, "intitule": m.normaliser_intitule,
        }
        for cas in lire("normaliseurs.attendu.json"):
            for cle, fonction in fonctions.items():
                with self.subTest(brut=cas["brut"], fonction=cle):
                    self.assertEqual(ecarts(cas[cle], fonction(cas["brut"])), [])

    def test_mappings(self):
        for cas in lire("mappings.attendu.json"):
            self.assertEqual(m.deviner_mapping(cas["entetes"], m.ENTETES_CANDIDATS), cas["candidats"])
            self.assertEqual(m.deviner_mapping(cas["entetes"], m.ENTETES_BLACKLIST), cas["blacklist"])

    def test_strategies(self):
        for cas in lire("strategies.attendu.json"):
            strategie = m.migrer_strategie(cas["strategie"])
            self.assertEqual(m.erreurs_strategie(strategie), cas["erreurs"])
            self.assertEqual(m.empreinte_strategie(strategie), cas["empreinte"])

    def test_empreintes(self):
        for cas in lire("empreintes.attendu.json"):
            self.assertEqual(m.empreinte_texte(cas["valeur"]), cas["empreinte"], cas["valeur"])

    def test_personas(self):
        attendu = lire("personas.attendu.json")
        for nom in ("defaut", "generique"):
            cas = attendu[nom]
            self.assertEqual(m.segmenter(cas["intitules"], cas["matrice"], False)["lignes"], cas["strict"], nom)
            self.assertEqual(m.segmenter(cas["intitules"], cas["matrice"], True)["lignes"], cas["multi"], nom)
            self.assertEqual([m.classer(i, cas["matrice"]["axes"][0]) for i in cas["intitules"]], cas["axe0"], nom)
        for cas in attendu["colonnes"]:
            self.assertEqual(m.deviner_colonne_poste(cas["entetes"]), cas["index"], cas["entetes"])


class LeBancSaitDireNon(unittest.TestCase):
    """Contrôles négatifs : un écart volontaire doit être vu."""

    def test_un_verdict_modifie_est_vu(self):
        attendu = lire("b2b-esn-validee.attendu.json")
        sabote = copy.deepcopy(attendu)
        sabote["lignes"][0]["statut"] = "Rejeté"
        self.assertTrue(ecarts(attendu, sabote))
        self.assertTrue(ecarts({"n": 20}, {"n": "20"}), "une chaîne ne vaut pas un nombre")

    def test_un_moteur_sabote_est_vu(self):
        original = m.email_est_professionnel
        m.email_est_professionnel = lambda valeur: True
        try:
            differences = ecarts(lire("b2b-esn-validee.attendu.json"), rendu_b2b(lire("b2b-esn-validee.entree.json")))
        finally:
            m.email_est_professionnel = original
        self.assertTrue(differences, "le banc n'a pas vu un moteur qui accepte tous les emails")


def ecrire_csv(chemin: Path, lignes: list[dict], separateur: str = ";") -> None:
    entetes = list(lignes[0]) if lignes else []
    with chemin.open("w", encoding="utf-8-sig", newline="") as fichier:
        ecrivain = csv.writer(fichier, delimiter=separateur, lineterminator="\n")
        ecrivain.writerow(entetes)
        for ligne in lignes:
            ecrivain.writerow(["" if ligne.get(e) is None else m.texte_js(ligne.get(e)) for e in entetes])


def lancer(*arguments: str) -> tuple[int, str]:
    sortie = io.StringIO()
    with contextlib.redirect_stdout(sortie), contextlib.redirect_stderr(io.StringIO()):
        code = m.main(list(arguments))
    return code, sortie.getvalue()


def preparer_mission(scenario: str) -> tuple[Path, Path, Path, Path]:
    """Un dossier temporaire avec l'export, la liste d'exclusion et la mission d'un scénario de parité."""
    dossier = Path(tempfile.mkdtemp(prefix="leadlint-"))
    entree = lire(f"b2b-{scenario}.entree.json")
    export, exclusions, mission = dossier / "export.csv", dossier / "exclusions.csv", dossier / "mission.json"
    ecrire_csv(export, entree["candidats"])
    ecrire_csv(exclusions, entree["blacklist"], ",")
    mission.write_text(json.dumps({"client": "Client fictif", "mission": "ESN 11-50", "configuration": entree["configuration"]}, ensure_ascii=False), encoding="utf-8")
    return dossier, export, exclusions, mission


class LigneDeCommande(unittest.TestCase):
    """Le script tel qu'un humain ou un agent le lance : fichiers réels, codes de sortie, classeur."""

    def preparer(self, scenario: str) -> tuple[Path, Path, Path, Path]:
        return preparer_mission(scenario)

    def test_mission_validee_rend_les_verdicts_de_la_page(self):
        dossier, export, exclusions, mission = self.preparer("esn-validee")
        code, sortie = lancer("b2b", str(export), "--moulinette", str(mission), "--blacklist", str(exclusions), "--json", "--lignes")
        self.assertEqual(code, 0, sortie)
        bilan = json.loads(sortie)
        attendu = lire("b2b-esn-validee.attendu.json")
        self.assertTrue(bilan["prevol"]["autorise"])
        self.assertEqual([l["statut"] for l in bilan["lignes"]], [l["statut"] for l in attendu["lignes"]])
        self.assertEqual([l["cause"] for l in bilan["lignes"]], [l["typeProbleme"] for l in attendu["lignes"]])
        self.assertEqual(bilan["total"], bilan["conformes"] + bilan["a_verifier"] + bilan["rejetes"])
        self.assertTrue(all(bilan["invariants"].values()))
        self.assertEqual(set(bilan["fichiers"]), {"controle_qa", "conformes"})

        import openpyxl
        classeur = openpyxl.load_workbook(bilan["fichiers"]["controle_qa"])
        self.assertEqual(classeur.sheetnames, ["Contrôle QA", "Audit moulinette", "Résumé"])
        controle = classeur["Contrôle QA"]
        self.assertEqual([c.value for c in controle[1]], m.ENTETES_CONTROLE_QA)
        self.assertEqual(controle.freeze_panes, "A2")
        self.assertIn("ControleQA", controle.tables)
        self.assertEqual(controle.max_row, len(attendu["lignes"]) + 1)
        self.assertTrue(controle.data_validations.dataValidation)
        conformes = openpyxl.load_workbook(bilan["fichiers"]["conformes"])["Conformes"]
        self.assertEqual(conformes.max_row - 1, bilan["conformes"])

    def test_excel_en_entree_donne_les_memes_verdicts(self):
        dossier, export, exclusions, mission = self.preparer("esn-validee")
        import openpyxl
        entree = lire("b2b-esn-validee.entree.json")
        classeur = openpyxl.Workbook()
        entetes = list(entree["candidats"][0])
        classeur.active.append(entetes)
        for ligne in entree["candidats"]:
            classeur.active.append([ligne.get(e) for e in entetes])
        xlsx = dossier / "export.xlsx"
        classeur.save(xlsx)
        code, sortie = lancer("b2b", str(xlsx), "--moulinette", str(mission), "--blacklist", str(exclusions), "--json", "--lignes")
        self.assertEqual(code, 0, sortie)
        self.assertEqual([l["statut"] for l in json.loads(sortie)["lignes"]], [l["statut"] for l in lire("b2b-esn-validee.attendu.json")["lignes"]])

    def test_prevol_bloque_ne_livre_rien(self):
        dossier, export, exclusions, mission = self.preparer("esn-non-validee")
        code, sortie = lancer("b2b", str(export), "--moulinette", str(mission), "--blacklist", str(exclusions), "--json")
        self.assertEqual(code, 2)
        self.assertFalse(json.loads(sortie)["livre"])
        self.assertEqual(sorted(p.name for p in dossier.glob("*.xlsx")), [])

    def test_diagnostic_rend_les_verdicts_sans_rien_de_livrable(self):
        dossier, export, exclusions, mission = self.preparer("esn-non-validee")
        code, sortie = lancer("b2b", str(export), "--moulinette", str(mission), "--blacklist", str(exclusions), "--json", "--diagnostic")
        self.assertEqual(code, 0, sortie)
        bilan = json.loads(sortie)
        self.assertTrue(bilan["diagnostic"])
        self.assertEqual(list(bilan["fichiers"]), ["controle_qa"])
        self.assertIn("-diagnostic-", bilan["fichiers"]["controle_qa"])

    def test_commande_prevol_et_empreintes(self):
        dossier, export, exclusions, mission = self.preparer("esn-validee")
        self.assertEqual(lancer("prevol", str(export), "--moulinette", str(mission), "--blacklist", str(exclusions))[0], 0)
        code, sortie = lancer("empreintes", "--moulinette", str(mission), "--blacklist", str(exclusions), "--json")
        self.assertEqual(code, 0)
        attendu = lire("b2b-esn-validee.attendu.json")["empreintes"]
        obtenu = json.loads(sortie)
        self.assertEqual({k: obtenu[k] for k in ("strategie", "recherche", "filtres", "blacklist")}, attendu)

    def test_une_formule_injectee_est_neutralisee(self):
        dossier, export, exclusions, mission = self.preparer("esn-validee")
        entree = lire("b2b-esn-validee.entree.json")
        entree["candidats"][0]["Nom complet"] = '=HYPERLINK("https://piege.example")'
        ecrire_csv(export, entree["candidats"])
        code, sortie = lancer("b2b", str(export), "--moulinette", str(mission), "--blacklist", str(exclusions), "--json", "--diagnostic")
        self.assertEqual(code, 0, sortie)
        import openpyxl
        cellule = openpyxl.load_workbook(json.loads(sortie)["fichiers"]["controle_qa"])["Contrôle QA"]["B2"]
        self.assertTrue(str(cellule.value).startswith("'="), cellule.value)

    def test_entree_illisible_sort_proprement(self):
        code, _ = lancer("b2b", "introuvable.csv", "--moulinette", "introuvable.json")
        self.assertEqual(code, 1)

    def test_personas(self):
        dossier = Path(tempfile.mkdtemp(prefix="leadlint-"))
        fichier = dossier / "contacts.csv"
        ecrire_csv(fichier, [{"Nom": "A", "Job title": "Directeur RSE"}, {"Nom": "B", "Job title": "Chargé IT"}, {"Nom": "C", "Job title": "Office manager"}], ",")
        code, sortie = lancer("personas", str(fichier), "--json")
        self.assertEqual(code, 0, sortie)
        bilan = json.loads(sortie)
        self.assertEqual(bilan["colonne"], "Job title")
        self.assertEqual(bilan["total"], 3)
        self.assertTrue(Path(bilan["fichier"]).exists())


class MissionTelechargee(unittest.TestCase):
    """Le pont web -> script : le fichier que le bouton « Télécharger la mission pour votre IA » fait télécharger passe le script.

    `mission-telechargee.json` est la sortie de `missionPourFichier` (oracle, scénario 6), validée dans la page sur l'export et la
    liste d'exclusion d'exemple de `skills/leadlint/exemples/` ; `mission-telechargee-attendu.json` porte les empreintes que la page
    calcule pour elle. Si la page change ce qu'elle écrit, ou le script ce qu'il lit, l'un de ces trois tests tombe.
    """

    EXPORT = EXEMPLES / "export-exemple.csv"
    EXCLUSIONS = EXEMPLES / "exclusions-exemple.csv"
    MISSION = PARITE / "mission-telechargee.json"

    def prevol(self, mission: Path) -> tuple[int, dict]:
        code, sortie = lancer("preflight", str(self.EXPORT), "--mission", str(mission), "--blacklist", str(self.EXCLUSIONS), "--json")
        return code, json.loads(sortie)["prevol"]

    def test_le_prevol_laisse_passer_la_mission_telechargee(self):
        code, prevol = self.prevol(self.MISSION)
        controles = prevol["controles"]
        self.assertEqual(code, 0, [c["id"] for c in controles if not c["valide"]])
        self.assertTrue(prevol["autorise"])
        self.assertEqual((sum(c["valide"] for c in controles), len(controles)), (11, 11))

    def test_les_empreintes_du_script_sont_celles_de_la_page(self):
        attendu = lire("mission-telechargee-attendu.json")
        self.assertEqual(set(attendu), {"strategie", "recherche", "filtres", "blacklist", "blacklist_volume"}, "un attendu amputé ne prouverait rien")
        code, sortie = lancer("fingerprints", "--mission", str(self.MISSION), "--blacklist", str(self.EXCLUSIONS), "--json")
        self.assertEqual(code, 0, sortie)
        obtenu = json.loads(sortie)
        self.assertEqual(ecarts(attendu, {champ: obtenu[champ] for champ in attendu if champ in obtenu}), [])

    def test_une_valeur_de_la_strategie_modifiee_apres_validation_bloque_le_prevol(self):
        """Contrôle négatif : la même mission, une seule valeur de sa stratégie changée en mémoire, n'est plus laissée passer."""
        self.assertEqual(self.prevol(self.MISSION)[0], 0, "le témoin positif : la mission d'origine passe, avec les mêmes fichiers")
        mission = json.loads(self.MISSION.read_text(encoding="utf-8"))
        mission["configuration"]["strategie"]["lignes"][0]["parametrage"] += " OU 50-99"
        with tempfile.TemporaryDirectory(prefix="leadlint-") as dossier:
            modifiee = Path(dossier) / "mission-modifiee.json"
            modifiee.write_text(json.dumps(mission, ensure_ascii=False), encoding="utf-8")
            code, prevol = self.prevol(modifiee)
        refuses = {c["id"] for c in prevol["controles"] if c["bloquant"] and not c["valide"]}
        self.assertEqual(code, 2)
        self.assertFalse(prevol["autorise"])
        self.assertIn("strategie", refuses)
        self.assertLessEqual(refuses, {"strategie", "recherche_validee"}, "seul le verrou de la stratégie (et la recherche qui l'embarque) doit tomber")


def cellules_du_classeur(chemin: str) -> dict:
    """Toutes les cellules d'un classeur, sauf l'horodatage du contrôle (la seule valeur qui change d'un passage à l'autre)."""
    import openpyxl
    rendu = {}
    for feuille in openpyxl.load_workbook(chemin).worksheets:
        for ligne in feuille.iter_rows():
            for cellule in ligne:
                if feuille.title == "Résumé" and cellule.column == 2 and feuille.cell(cellule.row, 1).value == "Contrôle effectué le":
                    continue
                rendu[(feuille.title, cellule.coordinate)] = cellule.value
    return rendu


class NomsAnglais(unittest.TestCase):
    """Les noms anglais sont des alias exacts des noms français : même sortie, octet pour octet, même code de sortie."""

    def deux_passages(self, francais: str, anglais: str, scenario: str, *options: str, avec_export: bool = True):
        """Le même appel sous le nom français et le nom anglais, dans le même dossier (donc avec les mêmes chemins écrits)."""
        dossier, export, exclusions, mission = preparer_mission(scenario)
        base = ([str(export)] if avec_export else []) + ["--blacklist", str(exclusions), *options]
        return lancer(francais, *base, "--moulinette", str(mission)), lancer(anglais, *base, "--mission", str(mission))

    def test_check_vaut_b2b(self):
        for scenario, options, code_attendu in (
            ("esn-validee", ("--json", "--lignes"), 0),
            ("esn-validee", (), 0),
            ("esn-non-validee", ("--json",), 2),
            ("esn-non-validee", (), 2),
            ("esn-non-validee", ("--json", "--diagnostic"), 0),
            ("esn-non-validee", ("--diagnostic",), 0),
        ):
            with self.subTest(scenario=scenario, options=options):
                (code_fr, sortie_fr), (code_en, sortie_en) = self.deux_passages("b2b", "check", scenario, *options)
                self.assertEqual((code_fr, code_en), (code_attendu, code_attendu))
                self.assertEqual(sortie_fr, sortie_en)
                self.assertTrue(sortie_fr)

    def test_check_ecrit_les_memes_classeurs_que_b2b(self):
        dossier, export, exclusions, mission = preparer_mission("esn-validee")
        base = [str(export), "--moulinette", str(mission), "--blacklist", str(exclusions), "--json"]
        code_fr, sortie_fr = lancer("b2b", *base)
        fichiers_fr = {cle: cellules_du_classeur(chemin) for cle, chemin in json.loads(sortie_fr)["fichiers"].items()}
        code_en, sortie_en = lancer("check", *base)
        fichiers_en = {cle: cellules_du_classeur(chemin) for cle, chemin in json.loads(sortie_en)["fichiers"].items()}
        self.assertEqual((code_fr, code_en), (0, 0))
        self.assertEqual(fichiers_fr, fichiers_en)
        self.assertTrue(all(fichiers_en.values()))

    def test_preflight_vaut_prevol(self):
        for scenario, code_attendu in (("esn-validee", 0), ("esn-non-validee", 2)):
            for options in ((), ("--json",)):
                with self.subTest(scenario=scenario, options=options):
                    (code_fr, sortie_fr), (code_en, sortie_en) = self.deux_passages("prevol", "preflight", scenario, *options)
                    self.assertEqual((code_fr, code_en), (code_attendu, code_attendu))
                    self.assertEqual(sortie_fr, sortie_en)
                    self.assertTrue(sortie_fr)

    def test_fingerprints_vaut_empreintes(self):
        for options in ((), ("--json",)):
            with self.subTest(options=options):
                (code_fr, sortie_fr), (code_en, sortie_en) = self.deux_passages("empreintes", "fingerprints", "esn-validee", *options, avec_export=False)
                self.assertEqual((code_fr, code_en), (0, 0))
                self.assertEqual(sortie_fr, sortie_en)
                self.assertTrue(sortie_fr)

    def test_les_alias_se_melangent_sans_changer_le_resultat(self):
        """Nom anglais avec l'option française, nom français avec l'option anglaise : toujours la même sortie."""
        dossier, export, exclusions, mission = preparer_mission("esn-validee")
        commun = [str(export), "--blacklist", str(exclusions), "--json", "--lignes"]
        reference = lancer("b2b", *commun, "--moulinette", str(mission))
        for nom, drapeau in (("check", "--moulinette"), ("b2b", "--mission")):
            with self.subTest(commande=nom, option=drapeau):
                self.assertEqual(lancer(nom, *commun, drapeau, str(mission)), reference)

    def test_entree_illisible_sort_en_1_sous_les_deux_noms(self):
        self.assertEqual(lancer("b2b", "introuvable.csv", "--moulinette", "introuvable.json")[0], 1)
        self.assertEqual(lancer("check", "introuvable.csv", "--mission", "introuvable.json")[0], 1)

    def test_la_mission_reste_obligatoire_sous_tous_les_noms(self):
        for nom in ("check", "b2b", "preflight", "prevol"):
            with self.subTest(commande=nom), self.assertRaises(SystemExit) as arret:
                lancer(nom, "export.csv")
            self.assertEqual(arret.exception.code, 2)
        for nom in ("fingerprints", "empreintes"):
            with self.subTest(commande=nom), self.assertRaises(SystemExit) as arret:
                lancer(nom)
            self.assertEqual(arret.exception.code, 2)

    def test_l_outil_se_presente_sous_son_nouveau_nom(self):
        dossier, export, exclusions, mission = preparer_mission("esn-validee")
        contacts = dossier / "contacts.csv"
        ecrire_csv(contacts, [{"Nom": "A", "Job title": "Directeur RSE"}], ",")
        sorties = {
            "check": lancer("check", str(export), "--mission", str(mission), "--blacklist", str(exclusions), "--json")[1],
            "preflight": lancer("preflight", str(export), "--mission", str(mission), "--blacklist", str(exclusions), "--json")[1],
            "fingerprints": lancer("fingerprints", "--mission", str(mission), "--json")[1],
            "personas": lancer("personas", str(contacts), "--json")[1],
        }
        for nom, sortie in sorties.items():
            with self.subTest(commande=nom):
                bilan = json.loads(sortie)
                self.assertEqual(bilan["outil"], "leadlint")
                self.assertEqual(bilan["version"], m.VERSION)
        resume = cellules_du_classeur(json.loads(sorties["check"])["fichiers"]["controle_qa"])
        lignes_outil = [v for (feuille, coordonnee), v in resume.items() if feuille == "Résumé" and coordonnee.startswith("B")]
        self.assertIn(f"leadlint {m.VERSION}", lignes_outil)
        self.assertNotIn(f"moulinette.py {m.VERSION}", lignes_outil)

    def test_les_sept_noms_sont_dans_l_aide(self):
        aide = io.StringIO()
        with contextlib.redirect_stdout(aide), self.assertRaises(SystemExit):
            m.main(["--help"])
        for mot in ("personas", "check", "b2b", "preflight", "prevol", "fingerprints", "empreintes"):
            self.assertIn(mot, aide.getvalue())


class Paquet(unittest.TestCase):
    """Le moteur est une seule copie ; installé en paquet, il retrouve ses matrices à côté de lui."""

    def installer_comme_un_paquet(self, matrice_generique: dict | None = None) -> Path:
        """Reproduit la disposition de site-packages : le module et `leadlint_data/matrices/` côte à côte."""
        site = Path(tempfile.mkdtemp(prefix="leadlint-site-"))
        shutil.copy(SCRIPT, site / "leadlint.py")
        donnees = site / "leadlint_data" / "matrices"
        shutil.copytree(RACINE / "skills" / "leadlint" / "matrices", donnees)
        if matrice_generique is not None:
            (donnees / "matrice-generique.json").write_text(json.dumps(matrice_generique, ensure_ascii=False), encoding="utf-8")
        return site

    def personas(self, script: Path) -> dict:
        dossier = Path(tempfile.mkdtemp(prefix="leadlint-travail-"))
        contacts = dossier / "contacts.csv"
        ecrire_csv(contacts, [{"Nom": "A", "Job title": "Directeur RSE"}, {"Nom": "B", "Job title": "Chargé IT"}], ",")
        resultat = subprocess.run(
            [sys.executable, str(script), "personas", str(contacts), "-o", str(dossier / "qualifie.xlsx"), "--json"],
            capture_output=True, text=True, encoding="utf-8", cwd=dossier,
        )
        self.assertEqual(resultat.returncode, 0, resultat.stderr)
        bilan = json.loads(resultat.stdout)
        bilan.pop("fichier")
        return bilan

    def test_matrice_par_defaut_trouvee_depuis_la_competence(self):
        self.assertEqual(m.MATRICE_DEFAUT, RACINE / "skills" / "leadlint" / "matrices" / "matrice-generique.json")
        self.assertTrue(m.MATRICE_DEFAUT.is_file())

    def test_matrice_par_defaut_trouvee_une_fois_installe(self):
        site = self.installer_comme_un_paquet()
        self.assertEqual(self.personas(site / "leadlint.py"), self.personas(SCRIPT))

    def test_la_matrice_livree_a_cote_du_module_est_bien_celle_lue(self):
        """Contrôle négatif : une matrice marquée, posée à côté du module installé, se retrouve dans le résultat."""
        marquee = json.loads((RACINE / "skills" / "leadlint" / "matrices" / "matrice-generique.json").read_text(encoding="utf-8"))
        marquee["axes"][0]["libelle"] = "AXE-MARQUE-PAR-LE-TEST"
        site = self.installer_comme_un_paquet(marquee)
        self.assertEqual(self.personas(site / "leadlint.py")["axes"][0]["libelle"], "AXE-MARQUE-PAR-LE-TEST")
        self.assertNotEqual(self.personas(SCRIPT)["axes"][0]["libelle"], "AXE-MARQUE-PAR-LE-TEST")

    def test_un_seul_moteur_dans_le_depot(self):
        moteurs = [p for p in RACINE.rglob("leadlint.py") if not {".git", "dist", "build", ".venv"} & set(p.parts)]
        self.assertEqual(moteurs, [SCRIPT])

    def test_versions_alignees(self):
        for nom in (".claude-plugin/plugin.json", ".claude-plugin/marketplace.json"):
            donnees = json.loads((RACINE / nom).read_text(encoding="utf-8"))
            self.assertEqual(donnees.get("version") or donnees["metadata"]["version"], m.VERSION, nom)

    def test_pyproject_declare_la_commande_et_la_source_unique(self):
        try:
            import tomllib
        except ImportError:  # Python < 3.11
            self.skipTest("tomllib absent (Python < 3.11)")
        projet = tomllib.loads((RACINE / "pyproject.toml").read_text(encoding="utf-8"))
        self.assertEqual(projet["project"]["name"], "leadlint")
        self.assertEqual(projet["project"]["scripts"], {"leadlint": "leadlint:main"})
        self.assertIn("version", projet["project"]["dynamic"])
        self.assertEqual(projet["tool"]["hatch"]["version"]["path"], "skills/leadlint/scripts/leadlint.py")
        livre = projet["tool"]["hatch"]["build"]["targets"]["wheel"]["force-include"]
        self.assertEqual(livre["skills/leadlint/scripts/leadlint.py"], "leadlint.py")
        for source in livre:
            self.assertTrue((RACINE / source).exists(), source)


if __name__ == "__main__":
    unittest.main(verbosity=1)
