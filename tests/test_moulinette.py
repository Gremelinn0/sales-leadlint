# -*- coding: utf-8 -*-
"""Banc de la moulinette : parité avec le moteur de la page web, puis la ligne de commande de bout en bout.

    python tests/test_moulinette.py

Les attendus de `tests/parite/` sont rendus par le moteur de la page web (oracle
`scripts/segmentation/oracle_parite.ts` du dépôt de la Marketplace). Le banc rejoue les
mêmes entrées dans le script et exige le même résultat, champ par champ. Deux contrôles
sont négatifs : un banc incapable de dire non ne vaut pas son oui.
"""

from __future__ import annotations

import copy
import contextlib
import csv
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path

RACINE = Path(__file__).resolve().parents[1]
PARITE = RACINE / "tests" / "parite"
sys.path.insert(0, str(RACINE / "skills" / "moulinette-b2b" / "scripts"))

import moulinette as m  # noqa: E402


def lire(nom: str):
    return json.loads((PARITE / nom).read_text(encoding="utf-8"))


def rendu_b2b(entree: dict) -> dict:
    """Ce que le script rend pour une entrée de parité, au format de l'oracle."""
    configuration = m.migrer_configuration_b2b(copy.deepcopy(entree["configuration"]))
    candidats, blacklist = entree["candidats"], entree["blacklist"]
    empreinte_bl = m.empreinte_blacklist(blacklist, configuration["mappingBlacklist"]) if blacklist else None
    controles = m.controler_prevol(configuration, len(candidats), len(blacklist), empreinte_bl)
    lignes = m.valider_export_b2b(candidats, blacklist, configuration)
    return {
        "empreintes": {
            "strategie": m.empreinte_strategie(configuration["strategie"]), "recherche": m.empreinte_recherche(configuration),
            "filtres": m.empreinte_filtres(configuration["pharow"]["filtres"]), "blacklist": empreinte_bl,
        },
        "prevol": {"autorise": m.prevol_autorise(controles), "controles": controles},
        "lignes": [{
            "numeroLigne": l["numeroLigne"], "statut": l["statut"], "cause": l["causePrincipale"]["critereId"] if l["causePrincipale"] else None,
            "typeProbleme": m.type_probleme(l), "raisons": l["raisons"], "valeurs": l["valeurs"], "evaluations": l["evaluations"],
        } for l in lignes],
    }


def ecarts(attendu, obtenu, chemin="") -> list[str]:
    """Toutes les différences, avec leur chemin ; un nombre vaut l'autre s'il est égal (20 == 20.0)."""
    if isinstance(attendu, dict) and isinstance(obtenu, dict):
        sortie = []
        for cle in sorted(set(attendu) | set(obtenu)):
            if cle not in obtenu or cle not in attendu:
                sortie.append(f"{chemin}.{cle} : présent d'un seul côté")
            else:
                sortie += ecarts(attendu[cle], obtenu[cle], f"{chemin}.{cle}")
        return sortie
    if isinstance(attendu, list) and isinstance(obtenu, list):
        if len(attendu) != len(obtenu):
            return [f"{chemin} : {len(attendu)} éléments attendus, {len(obtenu)} obtenus"]
        return [e for i, (a, o) in enumerate(zip(attendu, obtenu)) for e in ecarts(a, o, f"{chemin}[{i}]")]
    if m.est_nombre(attendu) and m.est_nombre(obtenu):
        return [] if attendu == obtenu else [f"{chemin} : {attendu!r} attendu, {obtenu!r} obtenu"]
    return [] if attendu == obtenu and type(attendu) is type(obtenu) else [f"{chemin} : {attendu!r} attendu, {obtenu!r} obtenu"]


SCENARIOS_B2B = ["esn-validee", "esn-non-validee", "esn-filtre-modifie", "regles-variees"]


class Parite(unittest.TestCase):
    """Le script rend exactement ce que rend la page web."""

    def test_scenarios_b2b(self):
        for nom in SCENARIOS_B2B:
            with self.subTest(scenario=nom):
                differences = ecarts(lire(f"b2b-{nom}.attendu.json"), rendu_b2b(lire(f"b2b-{nom}.entree.json")))
                self.assertEqual(differences, [], "\n".join(differences[:15]))

    def test_mission_validee_passe_le_prevol_apres_rechargement(self):
        self.assertTrue(lire("b2b-esn-validee.attendu.json")["prevol"]["autorise"])
        self.assertTrue(rendu_b2b(lire("b2b-esn-validee.entree.json"))["prevol"]["autorise"])

    def test_ancienne_mission(self):
        entree, attendu = lire("b2b-ancienne.entree.json"), lire("b2b-ancienne.attendu.json")
        configuration = m.migrer_configuration_b2b(entree["configuration"])
        regles = [{k: r[k] for k in ("critereId", "operateur", "valeur", "importance", "siAbsent")} for r in configuration["regles"]]
        self.assertEqual(regles, attendu["regles"])
        self.assertEqual([l["statut"] for l in m.valider_export_b2b(entree["candidats"], [], configuration)], attendu["statuts"])

    def test_normaliseurs(self):
        fonctions = {
            "texte": m.normaliser_texte, "nomEntreprise": m.normaliser_nom_entreprise, "linkedin": m.normaliser_url_linkedin,
            "domaine": m.normaliser_domaine, "siren": m.normaliser_siren, "categorie": m.normaliser_categorie_linkedin,
            "nombre": m.lire_nombre, "email": m.normaliser_email, "emailPro": m.email_est_professionnel,
            "mobile": m.normaliser_mobile_francais, "intitule": m.normaliser_intitule,
        }
        for cas in lire("normaliseurs.attendu.json"):
            for cle, fonction in fonctions.items():
                with self.subTest(brut=cas["brut"], fonction=cle):
                    self.assertEqual(ecarts(cas[cle], fonction(cas["brut"])), [])

    def test_mappings(self):
        for cas in lire("mappings.attendu.json"):
            self.assertEqual(m.deviner_mapping(cas["entetes"], m.ENTETES_CANDIDATS), cas["candidats"])
            self.assertEqual(m.deviner_mapping(cas["entetes"], m.ENTETES_BLACKLIST), cas["blacklist"])

    def test_strategies(self):
        for cas in lire("strategies.attendu.json"):
            strategie = m.migrer_strategie(cas["strategie"])
            self.assertEqual(m.erreurs_strategie(strategie), cas["erreurs"])
            self.assertEqual(m.empreinte_strategie(strategie), cas["empreinte"])

    def test_empreintes(self):
        for cas in lire("empreintes.attendu.json"):
            self.assertEqual(m.empreinte_texte(cas["valeur"]), cas["empreinte"], cas["valeur"])

    def test_personas(self):
        attendu = lire("personas.attendu.json")
        for nom in ("defaut", "generique"):
            cas = attendu[nom]
            self.assertEqual(m.segmenter(cas["intitules"], cas["matrice"], False)["lignes"], cas["strict"], nom)
            self.assertEqual(m.segmenter(cas["intitules"], cas["matrice"], True)["lignes"], cas["multi"], nom)
            self.assertEqual([m.classer(i, cas["matrice"]["axes"][0]) for i in cas["intitules"]], cas["axe0"], nom)
        for cas in attendu["colonnes"]:
            self.assertEqual(m.deviner_colonne_poste(cas["entetes"]), cas["index"], cas["entetes"])


class LeBancSaitDireNon(unittest.TestCase):
    """Contrôles négatifs : un écart volontaire doit être vu."""

    def test_un_verdict_modifie_est_vu(self):
        attendu = lire("b2b-esn-validee.attendu.json")
        sabote = copy.deepcopy(attendu)
        sabote["lignes"][0]["statut"] = "Rejeté"
        self.assertTrue(ecarts(attendu, sabote))
        self.assertTrue(ecarts({"n": 20}, {"n": "20"}), "une chaîne ne vaut pas un nombre")

    def test_un_moteur_sabote_est_vu(self):
        original = m.email_est_professionnel
        m.email_est_professionnel = lambda valeur: True
        try:
            differences = ecarts(lire("b2b-esn-validee.attendu.json"), rendu_b2b(lire("b2b-esn-validee.entree.json")))
        finally:
            m.email_est_professionnel = original
        self.assertTrue(differences, "le banc n'a pas vu un moteur qui accepte tous les emails")


def ecrire_csv(chemin: Path, lignes: list[dict], separateur: str = ";") -> None:
    entetes = list(lignes[0]) if lignes else []
    with chemin.open("w", encoding="utf-8-sig", newline="") as fichier:
        ecrivain = csv.writer(fichier, delimiter=separateur, lineterminator="\n")
        ecrivain.writerow(entetes)
        for ligne in lignes:
            ecrivain.writerow(["" if ligne.get(e) is None else m.texte_js(ligne.get(e)) for e in entetes])


def lancer(*arguments: str) -> tuple[int, str]:
    sortie = io.StringIO()
    with contextlib.redirect_stdout(sortie), contextlib.redirect_stderr(io.StringIO()):
        code = m.main(list(arguments))
    return code, sortie.getvalue()


class LigneDeCommande(unittest.TestCase):
    """Le script tel qu'un humain ou un agent le lance : fichiers réels, codes de sortie, classeur."""

    def preparer(self, scenario: str) -> tuple[Path, Path, Path, Path]:
        dossier = Path(tempfile.mkdtemp(prefix="moulinette-"))
        entree = lire(f"b2b-{scenario}.entree.json")
        export, exclusions, mission = dossier / "export.csv", dossier / "exclusions.csv", dossier / "mission.json"
        ecrire_csv(export, entree["candidats"])
        ecrire_csv(exclusions, entree["blacklist"], ",")
        mission.write_text(json.dumps({"client": "Client fictif", "mission": "ESN 11-50", "configuration": entree["configuration"]}, ensure_ascii=False), encoding="utf-8")
        return dossier, export, exclusions, mission

    def test_mission_validee_rend_les_verdicts_de_la_page(self):
        dossier, export, exclusions, mission = self.preparer("esn-validee")
        code, sortie = lancer("b2b", str(export), "--moulinette", str(mission), "--blacklist", str(exclusions), "--json", "--lignes")
        self.assertEqual(code, 0, sortie)
        bilan = json.loads(sortie)
        attendu = lire("b2b-esn-validee.attendu.json")
        self.assertTrue(bilan["prevol"]["autorise"])
        self.assertEqual([l["statut"] for l in bilan["lignes"]], [l["statut"] for l in attendu["lignes"]])
        self.assertEqual([l["cause"] for l in bilan["lignes"]], [l["typeProbleme"] for l in attendu["lignes"]])
        self.assertEqual(bilan["total"], bilan["conformes"] + bilan["a_verifier"] + bilan["rejetes"])
        self.assertTrue(all(bilan["invariants"].values()))
        self.assertEqual(set(bilan["fichiers"]), {"controle_qa", "conformes"})

        import openpyxl
        classeur = openpyxl.load_workbook(bilan["fichiers"]["controle_qa"])
        self.assertEqual(classeur.sheetnames, ["Contrôle QA", "Audit moulinette", "Résumé"])
        controle = classeur["Contrôle QA"]
        self.assertEqual([c.value for c in controle[1]], m.ENTETES_CONTROLE_QA)
        self.assertEqual(controle.freeze_panes, "A2")
        self.assertIn("ControleQA", controle.tables)
        self.assertEqual(controle.max_row, len(attendu["lignes"]) + 1)
        self.assertTrue(controle.data_validations.dataValidation)
        conformes = openpyxl.load_workbook(bilan["fichiers"]["conformes"])["Conformes"]
        self.assertEqual(conformes.max_row - 1, bilan["conformes"])

    def test_excel_en_entree_donne_les_memes_verdicts(self):
        dossier, export, exclusions, mission = self.preparer("esn-validee")
        import openpyxl
        entree = lire("b2b-esn-validee.entree.json")
        classeur = openpyxl.Workbook()
        entetes = list(entree["candidats"][0])
        classeur.active.append(entetes)
        for ligne in entree["candidats"]:
            classeur.active.append([ligne.get(e) for e in entetes])
        xlsx = dossier / "export.xlsx"
        classeur.save(xlsx)
        code, sortie = lancer("b2b", str(xlsx), "--moulinette", str(mission), "--blacklist", str(exclusions), "--json", "--lignes")
        self.assertEqual(code, 0, sortie)
        self.assertEqual([l["statut"] for l in json.loads(sortie)["lignes"]], [l["statut"] for l in lire("b2b-esn-validee.attendu.json")["lignes"]])

    def test_prevol_bloque_ne_livre_rien(self):
        dossier, export, exclusions, mission = self.preparer("esn-non-validee")
        code, sortie = lancer("b2b", str(export), "--moulinette", str(mission), "--blacklist", str(exclusions), "--json")
        self.assertEqual(code, 2)
        self.assertFalse(json.loads(sortie)["livre"])
        self.assertEqual(sorted(p.name for p in dossier.glob("*.xlsx")), [])

    def test_diagnostic_rend_les_verdicts_sans_rien_de_livrable(self):
        dossier, export, exclusions, mission = self.preparer("esn-non-validee")
        code, sortie = lancer("b2b", str(export), "--moulinette", str(mission), "--blacklist", str(exclusions), "--json", "--diagnostic")
        self.assertEqual(code, 0, sortie)
        bilan = json.loads(sortie)
        self.assertTrue(bilan["diagnostic"])
        self.assertEqual(list(bilan["fichiers"]), ["controle_qa"])
        self.assertIn("-diagnostic-", bilan["fichiers"]["controle_qa"])

    def test_commande_prevol_et_empreintes(self):
        dossier, export, exclusions, mission = self.preparer("esn-validee")
        self.assertEqual(lancer("prevol", str(export), "--moulinette", str(mission), "--blacklist", str(exclusions))[0], 0)
        code, sortie = lancer("empreintes", "--moulinette", str(mission), "--blacklist", str(exclusions), "--json")
        self.assertEqual(code, 0)
        attendu = lire("b2b-esn-validee.attendu.json")["empreintes"]
        obtenu = json.loads(sortie)
        self.assertEqual({k: obtenu[k] for k in ("strategie", "recherche", "filtres", "blacklist")}, attendu)

    def test_une_formule_injectee_est_neutralisee(self):
        dossier, export, exclusions, mission = self.preparer("esn-validee")
        entree = lire("b2b-esn-validee.entree.json")
        entree["candidats"][0]["Nom complet"] = '=HYPERLINK("https://piege.example")'
        ecrire_csv(export, entree["candidats"])
        code, sortie = lancer("b2b", str(export), "--moulinette", str(mission), "--blacklist", str(exclusions), "--json", "--diagnostic")
        self.assertEqual(code, 0, sortie)
        import openpyxl
        cellule = openpyxl.load_workbook(json.loads(sortie)["fichiers"]["controle_qa"])["Contrôle QA"]["B2"]
        self.assertTrue(str(cellule.value).startswith("'="), cellule.value)

    def test_entree_illisible_sort_proprement(self):
        code, _ = lancer("b2b", "introuvable.csv", "--moulinette", "introuvable.json")
        self.assertEqual(code, 1)

    def test_personas(self):
        dossier = Path(tempfile.mkdtemp(prefix="moulinette-"))
        fichier = dossier / "contacts.csv"
        ecrire_csv(fichier, [{"Nom": "A", "Job title": "Directeur RSE"}, {"Nom": "B", "Job title": "Chargé IT"}, {"Nom": "C", "Job title": "Office manager"}], ",")
        code, sortie = lancer("personas", str(fichier), "--json")
        self.assertEqual(code, 0, sortie)
        bilan = json.loads(sortie)
        self.assertEqual(bilan["colonne"], "Job title")
        self.assertEqual(bilan["total"], 3)
        self.assertTrue(Path(bilan["fichier"]).exists())


if __name__ == "__main__":
    unittest.main(verbosity=1)
